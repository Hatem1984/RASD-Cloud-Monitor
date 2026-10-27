"""Read-only QA comparison: two DIFFERENT historical publisher-source days.

Consumes owner's existing historical Word (24 Oct) and Gate E case (17 Oct).
Creates LOCAL QA-only before/after Word. No network, Gmail, scheduler or CAS.
"""
from __future__ import annotations
import argparse,json,hashlib,re
from pathlib import Path
from docx import Document
from daily_report import render_docx
from editorial_quality import build_editorial_brief,EditorialHold

def sample(day,headline,url,publisher):
 return {
  "title":"مَآل | MA'AL — مراجعة جودة تحليلات رصد التاريخية ("+day+")",
  "display_window_start":day+" 00:00 (تاريخ أرشيفي)",
  "display_window_end":day+" 23:59 (تاريخ أرشيفي)",
  "coverage":"HEALTHY","verified_links":1,
  "preview_only":True,"window_article_count":1,"truncated":False,
  "editorial_archive_only":True,"editorial_allowed_hosts":
     ("www.mped.gov.eg","mped.gov.eg"),
  "sources":[],"limitations":
     "مقارنة من مصادر أرشيفية محفوظة لدى المالك. لا تحقق جديدًا من الموقع "
     "ولا إعادة نشر للتقرير ولا إثباتًا لتسليم البريد. لا تُختلق أخبار أو أرقام.",
  "articles":[{
    "id":"QA-ARCHIVE-"+day,"headline":headline,
    "publisher_excerpt":"","source":publisher,
    "url":url,"published_at_cairo":day+" (date only)",
    "published_display_cairo":day+" — يوم النشر فقط؛ الساعة غير منشورة"
  }]
 }

def extract_24_word(path):
 doc=Document(path)
 headlines=[p.text[3:].strip() for p in doc.paragraphs if p.text.startswith("1. ")]
 if len(headlines)!=1 or "5 مليارات" not in headlines[0]:
  raise ValueError("unsupported_original_archive_article_count_or_title")
 urls=set(str(r.target_ref) for r in doc.part.rels.values()
          if r.reltype.endswith("/hyperlink") and "mped.gov.eg/singlenews" in str(r.target_ref))
 if len(urls)!=1:
  raise ValueError("missing_or_ambiguous_original_word_publisher_hyperlink")
 url=urls.pop()
 if "id=7275" not in url:
  raise ValueError("unexpected_historical_ministry_reference")
 return sample("2026-10-24",headlines[0],url,"وزارة التخطيط — الموقع الرسمي")

def extract_17_gatee(path):
 raw=json.loads(path.read_text("utf8"))
 row=raw["evidence"]
 if (raw["type"]!="HISTORICAL_OFFICIAL_SOURCE_DEMONSTRATION" or
     row["published_day_cairo"]!="2026-10-17" or
     row["publication_precision"]!="DATE_ONLY" or
     "mped.gov.eg/singlenews?id=7269" not in row["url"]):
  raise ValueError("unexpected_gate_e_historical_provenance")
 return sample("2026-10-17",row["headline"],row["url"],row["publisher"])

def main():
 ap=argparse.ArgumentParser()
 ap.add_argument("--archive-docx",type=Path,required=True)
 ap.add_argument("--gate-e-json",type=Path,required=True)
 ap.add_argument("--out",type=Path,required=True)
 args=ap.parse_args()
 if not args.archive_docx.is_file() or not args.gate_e_json.is_file():raise ValueError("source_missing")
 args.out.mkdir(parents=True,exist_ok=True)
 summaries=[]
 for report in (extract_17_gatee(args.gate_e_json),extract_24_word(args.archive_docx)):
  item=report["articles"][0];day=item["id"][-10:]
  before=args.out/("before-"+day+".docx")
  after=args.out/("after-"+day+".docx")
  render_docx(report,before)
  render_docx(report,after,editorial_mode="qa_only")
  source=build_editorial_brief(
      headline=item["headline"],publisher_excerpt="",
      original_url=item["url"],publisher=item["source"],
      published_day=day,proof_status="LOCAL_ARCHIVE_QA",
      approved_hosts={"www.mped.gov.eg","mped.gov.eg"})
  old_text="\n".join(p.text for p in Document(before).paragraphs)
  new_text="\n".join(p.text for p in Document(after).paragraphs)
  required=("الحقيقة المثبتة في العنوان:","أهم الأرقام المنشورة:",
            "لماذا يهم الشركات؟","خطوة عملية مشروطة:","ما الذي لم يُثبت؟")
  if any(marker not in new_text for marker in required):
   raise RuntimeError("editorial_sections_missing")
  if "تحليل تحريري مرتبط بموضوع الخبر:" not in old_text:
   raise RuntimeError("legacy_baseline_missing")
  if "تحليل تحريري مرتبط بموضوع الخبر:" in new_text:
   raise RuntimeError("legacy_filler_leaked_into_quality_report")
  if day=="2026-10-24":
   assert "5 مليارات دولار" in new_text
   assert "5,000 شركة" not in new_text # absent from archive headline -> NEVER invent
  summaries.append({
   "day":day,"source_type":"OWNER_EXISTING_HISTORICAL_ARCHIVE",
   "publication_precision":"DATE_ONLY",
   "qa_only":True,"verified_original_rechecked_now":False,
   "headline":item["headline"],"source_url":item["url"],
   "sourced_numeric_spans":[x["exact_source_span"] for x in source["figures_from_exact_source_spans"]],
   "new_specific_editorial_sections":len(required),"legacy_single_generic_paragraph":True,
   "legacy_docx":str(before),"new_docx":str(after),
   "before_sha256":hashlib.sha256(before.read_bytes()).hexdigest(),
   "after_sha256":hashlib.sha256(after.read_bytes()).hexdigest(),
   "email_sent":False
  })
 manifest={"schema":"RASD_EDITORIAL_QA_COMPARISON_V1","brand":"مَآل | MA'AL",
            "distinct_completed_historical_days":len({x["day"] for x in summaries}),
            "demo_only":True,"no_mail_or_scheduling_changes":True,"results":summaries}
 (args.out/"editorial_before_after_manifest.json").write_text(
     json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf8")
 print(json.dumps({"qa_pass":True,"days":[x["day"] for x in summaries],
                   "new_sections_per_article":len(required),
                   "source_images_not_revalidated":True,"email_sent":False},ensure_ascii=False),flush=True)
if __name__=="__main__":main()
