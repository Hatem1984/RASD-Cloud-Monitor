"""RASD daily read-only Arabic Word preview of validated original RSS headlines.

Manual cloud report only. Never sends Gmail, reads private credentials or
mutates persistent state. RSS headline coverage is not comprehensive news coverage.
"""
from __future__ import annotations

import argparse
import html
import json
from datetime import datetime, timedelta, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from gate11_runner import candidate_from_feed
from monitor import inspect, read_feed

CAIRO = ZoneInfo("Africa/Cairo")
WINDOW_HOURS = 24
MAX_ARTICLES = 15
MAX_PREVIEW_ARTICLES = 20  # Preview-only; strict daily email cap stays 15.
ARABIC_MONTHS = (
    "يناير", "فبراير", "مارس", "أبريل", "مايو", "يونيو",
    "يوليو", "أغسطس", "سبتمبر", "أكتوبر", "نوفمبر", "ديسمبر",
)
SOURCE_LABELS = {
    "Hapi native news RSS": "جريدة حابي",
    "Al Masry Al Youm native news RSS": "المصري اليوم",
    "Al Shorouk native Economy RSS": "بوابة الشروق",
}
ARABIC_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")


def editorial_caution(headline: str) -> str:
    """Conservative, headline-dependent interpretation, never a fabricated fact."""
    if "يشيد" in headline or "مؤشرات" in headline:
        return ("التغطية تنقل تقييمًا أو تصريحًا إيجابيًا عن مؤشرات اقتصادية؛"
                " لا يعادل ذلك قياسًا مستقلاً ولا يثبت أثر سياسات محددة."
                " يلزم فحص البيانات الأصلية والمنهجية قبل استخلاص نتائج.")
    if "صندوق" in headline or "تمويل" in headline:
        return ("الخبر يتناول مناقشة أو مقترحًا مرتبطًا بالتمويل؛"
                " البحث أو الإعلان لا يثبت إنشاء الصندوق أو اعتماد المخصصات أو الصرف الفعلي."
                " ينبغي متابعة القرارات والاتفاقيات المنشورة رسميًا.")
    return ("هذه تغطية صحفية لتصريحات أو مباحثات؛ لا يكفي ورودها وحده"
            " لإثبات صدور قرارات تنفيذية أو تحقق نتائج قابلة للقياس."
            " ويستلزم التقييم متابعة المصادر الرسمية وتطور الحدث.")


def cairo_display_time(moment: datetime) -> str:
    """Human-readable fully Arabic dates avoid RTL reordering of ISO/EEST text."""
    local = moment.astimezone(CAIRO)
    value = (f"{local.day} {ARABIC_MONTHS[local.month - 1]} {local.year}"
             f"، الساعة {local.hour:02}:{local.minute:02} بتوقيت القاهرة")
    return value.translate(ARABIC_DIGITS)


def gather(config: dict, *, fetch=read_feed, now: datetime | None = None,
           start_cairo: datetime | None = None,
           max_articles: int = MAX_ARTICLES) -> dict:
    """Collect verified RSS headlines, optionally within an exact Cairo-day window.

    A 20-item cap is permitted for the PUBLIC preview only. Production proof
    builders continue calling with the original 15-item default.
    """
    checked = datetime.now(timezone.utc) if now is None else now
    if checked.tzinfo is None or checked.utcoffset() is None:
        raise ValueError("naive_report_clock_forbidden")
    checked = checked.astimezone(timezone.utc)
    if type(max_articles) is not int or not 1 <= max_articles <= MAX_PREVIEW_ARTICLES:
        raise ValueError("report_article_limit_invalid")
    if start_cairo is None:
        since = checked - timedelta(hours=WINDOW_HOURS)
    else:
        if start_cairo.tzinfo is None or start_cairo.utcoffset() is None:
            raise ValueError("naive_pilot_start_forbidden")
        local_start = start_cairo.astimezone(CAIRO)
        local_end = checked.astimezone(CAIRO)
        since = local_start.astimezone(timezone.utc)
        if (local_start.date() != local_end.date() or since > checked or
                since < checked - timedelta(hours=WINDOW_HOURS)):
            raise ValueError("pilot_window_must_be_today_and_within_24h")
    # Collect enough validated candidates from BOTH sources before sorting.
    # monitor.inspect defaults to 15 for lightweight probes, which may
    # otherwise hide fresh stories from the second publisher.
    observed = inspect(config, fetch=fetch, now=checked, article_limit=200)
    articles = []
    for original in observed["matched_public_articles"]:
        article = candidate_from_feed(original, config, checked)
        if article is None:
            continue
        stamp = parsedate_to_datetime(article["published"]).astimezone(timezone.utc)
        # Reports never represent future-dated RSS items as recent published news.
        if not since <= stamp <= checked:
            continue
        articles.append({
            "id": article["id"],
            "headline": article["headline"],
            "publisher_excerpt": article.get("publisher_excerpt", ""),
            "source": article["source"],
            "url": article["url"],
            "published_at_cairo": stamp.astimezone(CAIRO).strftime("%Y-%m-%d %H:%M %Z"),
            "published_display_cairo": cairo_display_time(stamp),
            "_published_utc": stamp,
        })
    articles.sort(key=lambda x: x["_published_utc"], reverse=True)
    unique = {}
    for article in articles:
        unique.setdefault(article["id"], article)  # Same URL is one article.
    articles = list(unique.values())
    total_window_articles = len(articles)
    articles = articles[:max_articles]
    for article in articles:
        del article["_published_utc"]
    sources = [{
        "feed_id": item["feed_id"],
        "status": item["status"],
        "verified_links": item["verified_links"],
        "headline_matches": item["headline_matches"],
        "error_type": item.get("error") if item["status"] != "ok" else None,
    } for item in observed["sources"]]
    coverage = "DEGRADED" if any(s["status"] != "ok" for s in sources) else "HEALTHY"
    return {
        "title": "رصد | التقرير الإخباري اليومي",
        "window_start_cairo": since.astimezone(CAIRO).strftime("%Y-%m-%d %H:%M %Z"),
        "window_end_cairo": checked.astimezone(CAIRO).strftime("%Y-%m-%d %H:%M %Z"),
        "display_window_start": cairo_display_time(since),
        "display_window_end": cairo_display_time(checked),
        "coverage": coverage,
        "verified_links": observed["public_verified_publisher_links"],
        "rss_matches": observed["planning_headline_matches"],
        "window_article_count": total_window_articles,
        "article_limit": max_articles,
        "truncated": total_window_articles > max_articles,
        "articles": articles,
        "sources": sources,
        "limitations": (
            f"النتائج مبنية على عناوين قنوات RSS الأصلية لعدد {len(config['feeds'])} من المصادر المعتمدة في هذه التجربة فقط. "
            "قد تتأخر القنوات أو تغفل أخباراً؛ هذه ليست تغطية إعلامية شاملة. "
            f"الحد الأقصى للعرض {max_articles} مادة مطابقة، ولا يتم اختلاق أخبار عند عدم وجود نتائج. "
            "التغطيات التحريرية المتعددة قد تتعلق بحدث واحد؛ تصنيف الأحداث المستقلة يحتاج مراجعة تحريرية."
        ),
        "email_sent": False,
        "scheduler_enabled": False,
        "private_state_used": False,
    }


def _rtl(paragraph):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    p_pr.append(bidi)


def _word_hyperlink(paragraph, url, label="فتح الخبر في موقع الناشر"):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement("w:r")
    prop = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "086A80")
    prop.append(color)
    run.append(prop)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def render_docx(report: dict, destination: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from PIL import Image

    doc = Document()
    section = doc.sections[0]
    if any(row.get("evidence_image_path") for row in report.get("articles", [])):
        # Two-source original-photo bulletins formerly leaked the 2-line
        # methodological note to an almost empty third page. Tighter but
        # still comfortably printable margins preserve exact JPEG bytes.
        section.top_margin = Inches(0.52)
        section.bottom_margin = Inches(0.50)
        section.left_margin = Inches(0.67)
        section.right_margin = Inches(0.67)
        compact_evidence = True
    else:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        compact_evidence = False
    styles = doc.styles
    styles["Normal"].font.name = "DejaVu Sans"
    styles["Normal"].font.size = Pt(9.5 if compact_evidence else 10)
    styles["Normal"].font.color.rgb = RGBColor(36, 52, 69)
    styles["Title"].font.name = "DejaVu Sans"
    styles["Title"].font.size = Pt(15 if compact_evidence else 17)
    styles["Title"].font.color.rgb = RGBColor(8, 106, 128)

    p = doc.add_paragraph(style="Title")
    _rtl(p)
    p.add_run(report["title"])
    p = doc.add_paragraph()
    _rtl(p)
    p.add_run("الفترة: " + report["display_window_start"] + " إلى " + report["display_window_end"])
    p = doc.add_paragraph()
    _rtl(p)
    p.add_run(
        "حالة المصادر: " + report["coverage"] +
        " | روابط الناشرين المفحوصة: " + str(report["verified_links"]) +
        " | مواد مطابقة في الفترة: " + str(len(report["articles"]))
    )

    if report.get("preview_only"):
        notice = doc.add_paragraph()
        _rtl(notice)
        notice.add_run(
            "بروفة / مسودة فقط — لا تحتوي صورًا موثقة لصفحات الناشرين، "
            "ولا يجوز اعتبارها التقرير النهائي أو إثبات اكتمال الرصد."
        ).bold = True
        audit = doc.add_paragraph()
        _rtl(audit)
        audit.add_run(
            "المواد المطابقة داخل الفترة: " + str(report["window_article_count"]) +
            " | المعروض: " + str(len(report["articles"])) +
            " | هل تم الاختصار؟ " + ("نعم" if report["truncated"] else "لا") +
            " | الأحداث المستقلة: غير محددة آليًا"
        )
        counts = {}
        for row in report["articles"]:
            name = SOURCE_LABELS.get(row["source"], row["source"])
            counts[name] = counts.get(name, 0) + 1
        distribution = doc.add_paragraph()
        _rtl(distribution)
        distribution.add_run(
            "توزيع المواد حسب المصدر: " +
            ("؛ ".join(f"{name}: {count}" for name, count in counts.items())
             if counts else "لا توجد مواد مطابقة")
        )
    header = doc.add_paragraph()
    _rtl(header)
    header.add_run("الأخبار المطابقة").bold = True
    if not report["articles"]:
        no_items = doc.add_paragraph()
        _rtl(no_items)
        no_items.add_run("لم يُعثر على أخبار جديدة مطابقة في المصادر المتاحة خلال الفترة المحددة.")
    else:
        for index, article in enumerate(report["articles"], 1):
            p = doc.add_paragraph()
            _rtl(p)
            run = p.add_run(str(index) + ". " + article["headline"])
            run.bold = True
            p.paragraph_format.keep_with_next = True
            meta = doc.add_paragraph()
            _rtl(meta)
            meta.add_run("المصدر: " + SOURCE_LABELS.get(article["source"], article["source"]) +
                         " | النشر: " + article.get("published_display_cairo", article["published_at_cairo"]))
            meta.paragraph_format.keep_with_next = True
            link_paragraph = doc.add_paragraph()
            _rtl(link_paragraph)
            _word_hyperlink(link_paragraph, article["url"])
            link_paragraph.paragraph_format.keep_with_next = bool(article.get("publisher_excerpt"))
            if article.get("publisher_excerpt"):
                summary = doc.add_paragraph()
                _rtl(summary)
                summary.add_run("ملخص كما أورده الناشر في موجز RSS: ").bold = True
                summary.add_run(article["publisher_excerpt"])
                note = doc.add_paragraph()
                _rtl(note)
                note.add_run("قراءة تحليلية محدودة: " + editorial_caution(article["headline"]))
            if article.get("evidence_image_path"):
                evidence_note = doc.add_paragraph()
                _rtl(evidence_note)
                evidence_note.add_run("لقطة موثقة من صفحة الناشر الأصلية").italic = True
                # A genuine publisher full-page JPEG may be very tall. An
                # unconstrained 5.6-inch-wide image can exceed the printable
                # page height and become clipped in Word/LibreOffice. Preserve
                # aspect ratio and bound BOTH physical dimensions.
                with Image.open(article["evidence_image_path"]) as image:
                    width_px, height_px = image.size
                if width_px <= 0 or height_px <= 0:
                    raise ValueError("publisher_image_dimensions_invalid")
                aspect = height_px / width_px
                width_in = min(5.6, 6.2 / aspect)
                height_in = width_in * aspect
                doc.add_picture(
                    article["evidence_image_path"],
                    width=Inches(width_in), height=Inches(height_in))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    _rtl(p)
    p.add_run("حالة المصادر").bold = True
    for feed in report["sources"]:
        p = doc.add_paragraph()
        _rtl(p)
        p.add_run(
            {"hapi": "جريدة حابي", "almasryalyoum": "المصري اليوم"}.get(feed["feed_id"], feed["feed_id"]) +
            ": " + ("متاح" if feed["status"] == "ok" else "تعذّر الوصول") +
            " | روابط: " + str(feed["verified_links"]) +
            (" | فشل: " + feed["error_type"] if feed["error_type"] else "")
        )
    p = doc.add_paragraph()
    _rtl(p)
    p.add_run("تنبيه منهجي: " + report["limitations"])
    if compact_evidence:
        for paragraph in doc.paragraphs:
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.02
        # Keep the final scope notice intact on one page, rather than
        # isolating a fragment or changing meaning for aesthetics.
        doc.paragraphs[-1].paragraph_format.keep_together = True
    doc.core_properties.title = report["title"]
    doc.core_properties.author = "RASD"
    destination.parent.mkdir(parents=True, exist_ok=True)
    doc.save(destination)


def render_html(report: dict) -> str:
    esc = lambda val: html.escape(str(val), quote=True)
    def article_html(article, i):
        evidence = article.get("evidence_image_base64")
        screenshot_html = (
            '<figure><img class="evidence" alt="Publisher article screenshot" src="data:image/jpeg;base64,' +
            evidence + '"><figcaption>لقطة موثقة من صفحة الناشر الأصلية</figcaption></figure>'
            if evidence else ""
        )
        return (
            '<section class="article"><h2>' + str(i) + ". " + esc(article["headline"]) +
            '</h2><p class="meta">' + esc(article["source"]) + " | " +
            esc(article["published_at_cairo"]) +
            '</p><p class="url"><a rel="noopener noreferrer" href="' +
            esc(article["url"]) + '">' + esc(article["url"]) + "</a></p>" +
            ('<p class="summary"><strong>ملخص موجز الناشر:</strong> ' +
             esc(article["publisher_excerpt"]) + '</p>' if article.get("publisher_excerpt") else "") +
            screenshot_html + "</section>"
        )
    articles = "".join(article_html(article, i)
                       for i, article in enumerate(report["articles"], 1))
    if not articles:
        articles = '<p class="empty">لم يُعثر على أخبار جديدة مطابقة خلال آخر 24 ساعة.</p>'
    sources = "".join(
        "<li>" + esc(s["feed_id"]) + ": " + esc(s["status"]) +
        " (" + esc(s["verified_links"]) + " رابط)" +
        (" - " + esc(s["error_type"]) if s["error_type"] else "") + "</li>"
        for s in report["sources"]
    )
    status = "تغطية متدهورة: أحد المصادر غير متاح" if report["coverage"] == "DEGRADED" else "المصادر المتاحة سليمة"
    return (
        '<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8">'
        '<title>RASD Arabic daily report</title><style>'
        '@page{size:A4;margin:17mm 16mm;}'
        'html,body{margin:0;padding:0;background:white;color:#243445;}'
        'body{font-family:"DejaVu Sans","Noto Sans Arabic",sans-serif;font-size:12px;line-height:1.9;}'
        'header{border-bottom:3px solid #086a80;padding-bottom:12px;margin-bottom:18px;}'
        'h1{font-size:25px;color:#086a80;margin:0 0 8px;}'
        'h2{font-size:15px;color:#243445;line-height:1.6;margin:0 0 6px;}'
        '.meta{color:#536c7d;margin:0 0 4px;}'
        '.url{color:#086a80;font-size:9px;direction:ltr;text-align:left;overflow-wrap:anywhere;}'
        '.url a{color:inherit;text-decoration:underline;}'
        '.article{padding:10px 0;border-bottom:1px solid #dae7ec;break-inside:avoid;}'
        'figure{margin:12px 0 3px;text-align:center;break-inside:avoid;}'
        'img.evidence{display:block;max-width:100%;max-height:175mm;object-fit:contain;margin:auto;}'
        'figcaption{font-size:10px;color:#536c7d;margin-top:4px;}'
        '.badge{font-size:11px;color:#495f6b;}'
        '.empty{border:1px solid #cbdde3;border-radius:6px;padding:14px;}'
        'footer{border-top:1px solid #dae7ec;margin-top:22px;font-size:10px;color:#617582;}'
        '</style></head><body><header><h1>' + esc(report["title"]) +
        '</h1><div>' + esc(report["display_window_start"]) + " إلى " +
        esc(report["display_window_end"]) + '</div><div class="badge">' + esc(status) +
        " | روابط الناشرين: " + esc(report["verified_links"]) +
        " | الأخبار المدرجة: " + esc(len(report["articles"])) +
        "</div></header><main>" + articles +
        '<h2>حالة مصادر RSS</h2><ul>' + sources +
        '</ul></main><footer>' + esc(report["limitations"]) + '</footer></body></html>'
    )



def generate(config_path: Path, output_dir: Path, *, fetch=read_feed,
             now: datetime | None = None,
             start_cairo: datetime | None = None,
             max_articles: int = MAX_ARTICLES) -> dict:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    report = gather(config, fetch=fetch, now=now,
                    start_cairo=start_cairo, max_articles=max_articles)
    # Preview never claims to contain publisher-page screenshot evidence.
    # Only the strict cloud proof report may be treated as a final RASD report.
    report["title"] = "رصد | مسودة الفترة الصباحية — دون صور موثقة"
    report["preview_only"] = True
    stamp = report["window_end_cairo"][:10]
    docx = output_dir / ("RASD-Preview-" + stamp + ".docx")
    render_docx(report, docx)
    return {"mode": "manual_public_word_preview",
            "word_created": docx.is_file(),
            "preview_only": True,
            "articles": len(report["articles"]),
            "window_article_count": report["window_article_count"],
            "article_limit": report["article_limit"],
            "truncated": report["truncated"],
            "window_start_cairo": report["window_start_cairo"],
            "window_end_cairo": report["window_end_cairo"],
            "events_verified": None,
            "publisher_screenshots_verified": 0,
            "sources_healthy": sum(s["status"] == "ok" for s in report["sources"]),
            "sources_failed": sum(s["status"] != "ok" for s in report["sources"]),
            "coverage": report["coverage"],
            "email_sent": False,
            "scheduler_enabled": False,
            "private_state_used": False}

def main():
    parser = argparse.ArgumentParser(description="Generate source-backed Arabic Word preview (not final evidence)")
    parser.add_argument("--output-dir", type=Path, default=Path("report-output"))
    parser.add_argument("--start-cairo",
                        help="Exact Cairo-local start on today's date: YYYY-MM-DDTHH:MM")
    parser.add_argument("--max-articles", type=int, default=MAX_ARTICLES,
                        help="Public preview article cap, 1..20; default 15")
    args = parser.parse_args()
    try:
        start = (datetime.strptime(args.start_cairo, "%Y-%m-%dT%H:%M")
                 .replace(tzinfo=CAIRO) if args.start_cairo else None)
        result = generate(Path("sources.json"), args.output_dir,
                          start_cairo=start, max_articles=args.max_articles)
        print(json.dumps(result, ensure_ascii=False), flush=True)
        if result["coverage"] == "DEGRADED":
            print("::warning::RASD daily report source coverage DEGRADED", flush=True)
    except Exception:
        print("rasd_daily_report_failed_closed", flush=True)
        raise SystemExit(1) from None


if __name__ == "__main__":
    main()
