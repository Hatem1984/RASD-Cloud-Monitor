"""Local owner-only matching HTML/PDF/Word review exports from V0.3 JSON.

These are working review copies created on the owner's Windows device.
The independently page-verified 4-page Round2 Word/PDF are delivered directly
in ChatGPT; file copies generated here share source facts but may have different
pagination. No network, Gmail, CAS, model or production writes.
"""
from pathlib import Path
import json,html,hashlib
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from playwright.sync_api import sync_playwright
APP=Path(__file__).resolve().parent
data=json.loads((APP/"maal_web"/"product_report.json").read_text("utf8"))
assert data["signal_count"]==2 and len(data["stories"])==3
folder=Path(r"E:\MAAL\03_REPORTS_AND_REVIEW\MAAL_PRODUCT_ACCEPTANCE_R2_20261027\owner_downloads")
folder.mkdir(parents=True,exist_ok=True)
basename="MAAL_Business_Intelligence_Round2_2026-10-27"
word=folder/(basename+".docx")
htmlpath=folder/(basename+".html")
pdf=folder/(basename+".pdf")
if any(x.exists() for x in (word,htmlpath,pdf)):
 raise SystemExit("HOLD_EXPORT_ALREADY_EXISTS_DO_NOT_OVERWRITE")
D=Document();sec=D.sections[0]
sec.top_margin=Cm(1.7);sec.bottom_margin=Cm(1.7)
sec.left_margin=Cm(1.9);sec.right_margin=Cm(1.9)
for name,size in [("Normal",10.5),("Title",19),("Heading 1",14),("Heading 2",11.5)]:
 st=D.styles[name];st.font.name="Noto Naskh Arabic";st.font.size=Pt(size)
 st._element.get_or_add_rPr().rFonts.set(qn("w:cs"),"Noto Naskh Arabic")
def paragraph(text,style=None):
 p=D.add_paragraph(style=style);p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 pp=p._p.get_or_add_pPr();b=OxmlElement("w:bidi");b.set(qn("w:val"),"1");pp.append(b)
 p.add_run(text);return p
paragraph("مَآل | MA'AL — استخبارات أعمال", "Title")
paragraph("مراجعة الجولة الثانية | الفترة: "+data["period"]+" | ليست تغطية حية")
paragraph(data["headline_summary"])
paragraph(data["summary"])
paragraph("الشركة الافتراضية والافتراضات","Heading 1")
paragraph(data["company"]+". "+data["company_warning"])
paragraph("نطاق الرصد","Heading 1")
paragraph(data["coverage"]["note"]+". "+data["coverage"]["not_observed"])
for sig in data["signals"]:
 paragraph(sig["title"],"Heading 1")
 paragraph("الأهمية: "+sig["priority_label"]+" | القطاع: "+sig["sector_label"])
 paragraph("حالة الدليل: "+sig["verification_label"])
 for fact in sig["facts"]:paragraph("• "+fact)
 for label,key in [
 ("الأثر المحتمل","why"),("الفرضية المشروطة","hypothesis"),
 ("المخاطر","risk"),("البيانات الناقصة","missing"),
 ("التوصية","recommendation"),("المسؤول المقترح","suggested_owner"),
 ("المتابعة","follow_up")]: paragraph(label+": "+sig[key])
 for code in sig["source_ids"]:
  src=next(s for s in data["stories"] if s["id"]==code)
  paragraph("الناشر "+src["published_day"]+": "+src["source_url"]+" | "+src["evidence_status"])
paragraph("القرارات المقترحة — غير منفذة","Heading 1")
for rec in data["decisions"]:
 paragraph(str(rec["priority"])+". "+rec["title"]+" | "+rec["owner"])
 paragraph("شرط: "+rec["condition"]+" | متابعة: "+rec["follow_up"])
paragraph("حدود التحقق","Heading 1")
for line in data["limitations"]:paragraph("• "+line)
D.save(word)

def esc(text):return html.escape(str(text),quote=True)
blocks=[]
for sig in data["signals"]:
 inner="<h2>"+esc(sig["title"])+"</h2>"
 inner+="<p class='muted'>"+esc(sig["priority_label"]+" · "+sig["sector_label"]+" · "+sig["verification_label"])+"</p>"
 inner+="<h3>حقائق منشورة</h3><ul>"+"".join("<li>"+esc(x)+"</li>" for x in sig["facts"])+"</ul>"
 for label,key in [("لماذا يهم الشركة؟","why"),("الاستنتاج المشروط","hypothesis"),
                   ("المخاطر","risk"),("القرار المقترح","recommendation"),
                   ("المعلومات غير المثبتة","missing"),("المسؤول المقترح","suggested_owner"),
                   ("التوقيت","follow_up")]:
  inner+="<p><strong>"+label+": </strong>"+esc(sig[key])+"</p>"
 inner+="<h3>الأدلة وروابطها</h3>"
 for code in sig["source_ids"]:
  s=next(t for t in data["stories"] if t["id"]==code)
  inner+="<p class='source'>"+esc(s["published_day"]+" · "+s["evidence_status"])+"<br><a href='"+esc(s["source_url"])+"'>"+esc(s["source_url"])+"</a></p>"
 blocks.append("<section>"+inner+"</section>")
html_content=("""<!doctype html><html lang='ar' dir='rtl'><head><meta charset='utf-8'>
<style>@page{size:A4;margin:17mm}body{font-family:'Segoe UI',Tahoma,Arial,sans-serif;color:#143548;line-height:1.8;font-size:12px}
h1{font-size:23px}h2{font-size:17px;color:#0e696c}h3{font-size:13px;color:#2d6476}
header{border-bottom:3px solid #158780;margin-bottom:16px;padding-bottom:12px}
section{margin:16px 0;padding:8px 0;border-top:1px solid #c2d9dc;break-inside:avoid-page}
strong{color:#095c62}.muted{font-size:11px;color:#546a73}.source{font-size:10px;background:#ecf3f4;padding:9px;overflow-wrap:anywhere}
table{border-collapse:collapse;width:100%}td,th{border:1px solid #bbcfd4;padding:6px;text-align:right}a{color:#087778}
footer{color:#58737d;font-size:10px;border-top:1px solid #c9d7db;padding-top:8px}
</style></head><body><header><h1>مَآل | MA'AL — تقرير استخبارات أعمال</h1>
<p>مراجعة الجولة الثانية: """+esc(data["period"])+""" • تقرير تاريخي محدود؛ ليس رصدًا حيًا ولا إشعارًا بريديًا</p>
<p>"""+esc(data["headline_summary"])+"""</p></header>
<h2>الملخص التنفيذي</h2><p>"""+esc(data["summary"])+"""</p>
<h2>الشركة محل التحليل</h2><p>"""+esc(data["company"]+" — "+data["company_warning"])+"""</p>
<h2>نطاق البيانات</h2><p>"""+esc(data["coverage"]["note"]+" "+data["coverage"]["not_observed"])+"""</p>"""+
"".join(blocks)+"""<section><h2>قرارات مرتبة — للمراجعة فقط</h2>"""+
"".join("<p><strong>"+str(a["priority"])+". "+esc(a["title"])+"</strong><br>المسؤول: "+esc(a["owner"])+"<br>توقيت المتابعة: "+esc(a["follow_up"])+"<br>الشرط: "+esc(a["condition"])+"</p>" for a in data["decisions"])+
"""</section><h2>ما لم يُثبت</h2><ul>"""+
"".join("<li>"+esc(a)+"</li>" for a in data["limitations"])+"""</ul>
<footer>مَآل | MA'AL · مراجعة فقط · الشركة افتراضية · لم تُرسل بالبريد</footer></body></html>""")
htmlpath.write_text(html_content,encoding="utf8")
with sync_playwright() as p:
 browser=p.chromium.launch(headless=True)
 page=browser.new_page(locale="ar-EG")
 page.goto(htmlpath.as_uri(),wait_until="load")
 page.pdf(path=str(pdf),format="A4",print_background=True,prefer_css_page_size=True)
 browser.close()
assert word.stat().st_size>15000 and pdf.stat().st_size>15000
receipt={"word_sha256":hashlib.sha256(word.read_bytes()).hexdigest(),
 "pdf_sha256":hashlib.sha256(pdf.read_bytes()).hexdigest(),
 "source_json_sha256":hashlib.sha256((APP/"maal_web"/"product_report.json").read_bytes()).hexdigest(),
 "source_documents":3,"grouped_signals":2,"mail_sent":False,
 "print_pdf_generated_in_real_chromium":True,"word_generated_with_python_docx":True,
 "separate_from_chat_page_verified_4_page_review":True}
(folder/"owner_downloads_receipt.json").write_text(json.dumps(receipt,ensure_ascii=False,indent=2),encoding="utf8")
print("MAAL_V03_OWNER_REPORT_DOWNLOADS_CREATED",json.dumps({"pdf_bytes":pdf.stat().st_size,
 "word_bytes":word.stat().st_size,"signals":2},ensure_ascii=False),flush=True)
