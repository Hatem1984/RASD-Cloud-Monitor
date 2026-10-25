"""Documented October 25 production failures reproduced without SMTP or ledger."""
import tempfile
import unittest
from pathlib import Path

from docx import Document

from capture_public import is_planning_publisher_page
from daily_report import render_docx


class NextReportQualityGate(unittest.TestCase):
    def test_ministry_programme_without_literal_planning_is_valid(self):
        headline = 'مدبولي: حوارات مجتمعية حول البرنامج الوطني للتحول الاقتصادي'
        self.assertTrue(is_planning_publisher_page(headline, ' ' * 200))
        self.assertTrue(is_planning_publisher_page(
            'National Planning Institute discusses Egypt', ' ' * 200))
        self.assertTrue(is_planning_publisher_page(
            'وزارة التخطيط تبحث الخطة الاستثمارية للدولة', ' ' * 200))

    def test_generic_economy_and_sports_still_rejected(self):
        for headline in (
            'البرنامج الوطني لبطولة كرة القدم',
            'القطاع الخاص يناقش برنامج التمويل',
            'أسعار الدولار ترتفع في الأسواق',
            'Planning a football game',
        ):
            with self.subTest(headline=headline):
                self.assertFalse(is_planning_publisher_page(headline, ' ' * 200))

    def test_word_uses_real_rss_count_and_arabic_hold_reason(self):
        report = {
            'title': 'رصد | تقرير حالة المتابعة', 'coverage':
                'ORIGINAL_IMAGE_HOLD_NO_VERIFIED_DELIVERY',
            'display_window_start': '٢٥ أكتوبر الساعة ٢٣:٠٠',
            'display_window_end': '٢٥ أكتوبر الساعة ٢٣:٠٧',
            'verified_links': 218,
            'articles': [],
            'sources': [],
            'scope_note': '13 مصدرًا نشطًا ضمن سجل من 50 مصدرًا',
            'held_candidates': [{
                'headline': 'مدبولي: البرنامج الوطني للتحول الاقتصادي',
                'source': 'Alborsa native Egyptian Economy RSS',
                'url': 'https://www.alborsaanews.com/2026/10/22/2001319',
                'reason': 'not_a_planning_article',
                'published_display_cairo': '٢٢ أكتوبر ٢٠٢٦',
            }],
            'limitations': ('عدد الروابط لا يعني اكتمال سوق الأخبار،'
                            ' والمعلّق ليس خبرًا مُسلّمًا.'),
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'RASD-sample.docx'
            render_docx(report, path)
            doc = Document(path)
            body = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('218', body)
            self.assertIn('روابط RSS التي أمكن فحصها', body)
            self.assertIn('لا أخبار جديدة مستوفية؛ توجد مواد معلقة', body)
            self.assertIn('لم يتأكد ارتباط صفحة الناشر', body)
            self.assertIn('13 مصدرًا نشطًا', body)
            self.assertNotIn('لم يتم تجميع عدّادها', body)
            self.assertFalse(doc.inline_shapes)
            self.assertTrue(any(rel.target_ref == report['held_candidates'][0]['url']
                                for rel in doc.part.rels.values()
                                if rel.reltype.endswith('/hyperlink')))

    def test_daily_delivered_word_does_not_relabel_published_articles_as_today(self):
        report = {
            'title': 'رصد | التقرير المجمع للأخبار المرسلة | ٢٥ أكتوبر',
            'coverage': 'CONSOLIDATED_PREVIOUSLY_SENT_ORIGINALS',
            'display_window_start': '٢٤ أكتوبر، ٢٢:٠٠',
            'display_window_end': '٢٥ أكتوبر، ٢٢:٠٠',
            'verified_links': 0, 'articles': [], 'sources': [],
            'limitations': ('تاريخ نشر الخبر قد يسبق يوم إرسال التقرير؛ '
                            'هذه حصيلة تسليم وليست تغطية شاملة.'),
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'daily.docx'
            render_docx(report, path)
            body = '\n'.join(p.text for p in Document(path).paragraphs)
            self.assertIn('روابط الأخبار المسلّمة', body)
            self.assertIn('أخبار سبق تسليمها بصورها الأصلية', body)
            self.assertNotIn('روابط الناشرين المفحوصة: 0', body)


if __name__ == '__main__':
    unittest.main()
