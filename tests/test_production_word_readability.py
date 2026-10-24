"""Production Word editorial QA: readable Cairo window, all source statuses, original JPEG."""
import hashlib
import os
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document
from PIL import Image, ImageDraw

import scheduled_rasd as worker
from daily_report import render_docx


class ProductionWordQuality(unittest.TestCase):
    def test_word_retains_every_source_health_row_compactly(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            image_path = root / "publisher-test.jpg"
            picture = Image.new("RGB", (1200, 900), "white")
            draw = ImageDraw.Draw(picture)
            draw.rectangle((0, 0, 1200, 140), fill=(15, 90, 130))
            draw.text((70, 250), "OFFLINE TEST ONLY", fill=(0, 0, 0))
            picture.save(image_path)
            before = hashlib.sha256(image_path.read_bytes()).hexdigest()
            report = {
                "title": "رصد | اختبار تنسيقي فقط",
                "display_window_start": "٢٤ أكتوبر ٢٠٢٦، الساعة ١٤:٣٠ بتوقيت القاهرة",
                "display_window_end": "٢٤ أكتوبر ٢٠٢٦، الساعة ١٤:٥٥ بتوقيت القاهرة",
                "coverage": "SOURCE_LIMITED",
                "verified_links": 1,
                "articles": [{
                    "id": "a" * 24, "headline": "وزارة التخطيط تعلن مقترح إصلاح تشريعي",
                    "source": "Alborsa native Egyptian Economy RSS",
                    "url": "https://www.alborsaanews.com/2026/10/24/1234",
                    "published_at_cairo": "2026-10-24 13:00 EEST",
                    "published_display_cairo": "٢٤ أكتوبر ٢٠٢٦، الساعة ١٣:٠٠ بتوقيت القاهرة",
                    "publisher_excerpt": "مقترح مبدئي يخضع للمراجعة",
                    "evidence_image_path": str(image_path),
                }],
                "sources": [
                    {"feed_id": "hapi", "status": "failed", "verified_links": 0,
                     "error_type": "original_feed_unavailable"},
                    {"feed_id": "almasryalyoum", "status": "ok", "verified_links": 0,
                     "error_type": None},
                    {"feed_id": "shorouk", "status": "ok", "verified_links": 0,
                     "error_type": None},
                    {"feed_id": "alborsa", "status": "ok", "verified_links": 1,
                     "error_type": None},
                ],
                "limitations": "لا يُعد هذا الاختبار دليلًا على خبر حقيقي.",
            }
            output = root / "offline.docx"
            render_docx(report, output)
            doc = Document(output)
            self.assertEqual(len(doc.inline_shapes), 1)
            self.assertEqual(hashlib.sha256(image_path.read_bytes()).hexdigest(), before)
            all_text = "\n".join(p.text for p in doc.paragraphs)
            self.assertIn("٢٤ أكتوبر ٢٠٢٦", all_text)
            self.assertNotIn("2026-10-24T", all_text)
            self.assertEqual(sum("جريدة حابي:" in p.text for p in doc.paragraphs), 1)
            source_paragraphs = [p.text for p in doc.paragraphs if "جريدة حابي:" in p.text]
            self.assertEqual(len(source_paragraphs), 1)
            self.assertTrue(all(x in source_paragraphs[0] for x in
                                ("المصري اليوم", "بوابة الشروق", "جريدة البورصة")))
            self.assertIn("تعذّر الوصول", source_paragraphs[0])

    def test_production_interval_passes_Arabic_cairo_clock_to_word(self):
        ledger = MagicMock()
        ledger.read.return_value = ({"entries": {}}, "private-sha")
        story = {
            "id": "a" * 24,
            "headline": "وزارة التخطيط تستعرض برنامج التحول الاقتصادي",
            "source": "Alborsa native Egyptian Economy RSS",
            "url": "https://www.alborsaanews.com/2026/10/24/1234",
            "published_at_cairo": "2026-10-24 13:00 EEST",
        }
        observed = {}
        def build(items, config, output, **kwargs):
            observed.update(kwargs)
            word = output / "offline.docx"
            word.write_bytes(b"not real Word: this is only an offline mock")
            return word, items, 0, False
        def discover(_config, _now, *, allow_degraded, source_records):
            source_records.extend([
                {"feed_id": "hapi", "status": "ok"},
                {"feed_id": "almasryalyoum", "status": "ok"},
                {"feed_id": "shorouk", "status": "ok"},
                {"feed_id": "alborsa", "status": "ok"},
            ])
            return [story]
        with (
            patch.dict(os.environ, {"GITHUB_EVENT_NAME": "schedule",
                                   "RASD_JOB_CRON": worker.HALF_HOUR_CRON}),
            patch.object(worker, "require_github_scheduler"),
            patch.object(worker, "scheduled_phase", return_value="interval"),
            patch.object(worker, "protected_credentials",
                         return_value=("sender@gmail.com", "recipient@example.com",
                                       "offline-secret", ledger)),
            patch.object(worker, "discover_interval_news", side_effect=discover),
            patch.object(worker, "prepare_source_limited_verified_word", side_effect=build),
            patch.object(worker, "reserve_stories"),
            patch.object(worker, "send_one_word", return_value="<offline-no-smtp>"),
        ):
            result = worker.run_scheduled(
                "interval", datetime(2026, 10, 24, 13, 55, tzinfo=timezone.utc))
        self.assertEqual(result["items"], 1)
        self.assertIn("بتوقيت القاهرة", observed["start"])
        self.assertIn("بتوقيت القاهرة", observed["end"])
        self.assertNotIn("T", observed["start"])
        self.assertNotIn("T", observed["end"])


if __name__ == "__main__":
    unittest.main()
