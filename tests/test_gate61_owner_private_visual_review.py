"""Gate61: only the owner locally may attest to actual decrypted Word pages.

These are SYNTHETIC tests using an ephemeral private key; never publisher proof.
"""
from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo
from unittest.mock import patch

from docx import Document
from docx.shared import Inches

from gate43_owner_encrypted_review import pack_owner_review, unpack_owner_review
from gate61_owner_private_visual_review import (
    inspect_private_review, make_private_gallery, owner_local_attestation,
)
from test_gate43_owner_encrypted_review import _keys, audit
from test_gate50_owner_encrypted_review_pages import synthetic_private_page

DAY = date(2026, 10, 17)


def fixture(folder):
    root = Path(folder)
    pub, private = _keys(root)
    reports = root / "words"
    reports.mkdir()
    pages = root / "pages"
    png = synthetic_private_page(pages, 1)
    word = reports / "RASD-Unified-2026-10-17.docx"
    doc = Document()
    doc.add_paragraph("تقرير اصطناعي للاختبار، وليس خبرًا أصليًا")
    doc.add_picture(str(png), width=Inches(3.5))
    doc.save(word)
    facts = audit("READY_QA_ONLY")
    facts.update({
        "word_created": True,
        "official_pagination_complete": True,
        "source_health": [
            {"feed_id": "hapi", "status": "ok", "verified_links": 10},
            {"feed_id": "almasryalyoum", "status": "ok", "verified_links": 20},
            {"feed_id": "mped-official", "status": "ok", "verified_links": 20},
        ],
    })
    payload = root / "owner-ciphertext.json"
    pack_owner_review(reports, facts, payload, pub,
                      render_passed=True, review_pages_dir=pages)
    decrypted = root / "owner-private-decrypted"
    unpack_owner_review(payload, private, decrypted)
    return decrypted, word, payload


class LocalPrivateOwnerReviewTests(unittest.TestCase):
    def test_real_crypto_roundtrip_private_gallery_does_not_infer_human_approval(self):
        with tempfile.TemporaryDirectory() as temp:
            folder, original, _ = fixture(temp)
            details = inspect_private_review(folder, DAY)
            self.assertEqual(details["original_source_count"], 3)
            self.assertEqual(details["article_count"], 1)
            self.assertEqual(details["review_page_count"], 1)
            self.assertFalse(details["human_visual_review"])
            self.assertFalse(details["production_delivery_approved"])
            self.assertEqual(len(details["page_sha256"]), 1)
            gallery = make_private_gallery(folder, details)
            body = gallery.read_text(encoding="utf-8")
            self.assertIn("review-pages/page-001.png", body)
            self.assertIn(original.name, body)
            self.assertIn("تقرير", body)
            self.assertNotIn("<script", body)
            self.assertNotIn("http:", body)
            self.assertNotIn("https:", body)
            self.assertFalse((folder / "rasd-owner-human-review-attestation.json").exists())

    def test_current_or_future_cairo_day_can_never_receive_full_day_attestation(self):
        with tempfile.TemporaryDirectory() as temp:
            folder, _, _ = fixture(temp)
            incomplete = datetime.now(ZoneInfo("Africa/Cairo")).date()
            with self.assertRaisesRegex(ValueError, "incomplete_cairo_calendar_day"):
                inspect_private_review(folder, incomplete)

    def test_healthy_original_rss_with_zero_items_is_not_an_outage(self):
        # A well-formed empty publisher RSS is a healthy feed; the first-party
        # Ministry archive must still show a verifiable original index link.
        with tempfile.TemporaryDirectory() as temp:
            folder, _, _ = fixture(temp)
            path = folder / "rasd-owner-review-audit.json"
            data = json.loads(path.read_text(encoding="utf-8"))
            data["source_health"][0]["verified_links"] = 0
            data["source_health"][1]["verified_links"] = 0
            path.write_text(json.dumps(data), encoding="utf-8")
            verified = inspect_private_review(folder, DAY)
            self.assertEqual(verified["original_source_count"], 3)
            data["source_health"][2]["verified_links"] = 0
            path.write_text(json.dumps(data), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "original_sources_not_healthy"):
                inspect_private_review(folder, DAY)

    def test_wrong_day_and_missing_document_fail(self):
        with tempfile.TemporaryDirectory() as temp:
            folder, word, _ = fixture(temp)
            with self.assertRaisesRegex(ValueError, "private_review_audit_not_qualified"):
                inspect_private_review(folder, date(2026, 10, 16))
            (folder / word.name).unlink()
            with self.assertRaisesRegex(ValueError, "private_review_word_missing"):
                inspect_private_review(folder, DAY)

    def test_source_omission_prohibited_email_or_forged_upstream_human_review_fails(self):
        for field, value in (
            ("email_sent", True),
            ("human_visual_review", True),
            ("source_added_to_production", True),
            ("word_images_verified", False),
            ("owner_review_pages_encrypted", False),
            ("official_pagination_complete", False),
            ("target_date_cairo", "2026-10-18"),
        ):
            with self.subTest(field=field), tempfile.TemporaryDirectory() as temp:
                folder, _, _ = fixture(temp)
                path = folder / "rasd-owner-review-audit.json"
                data = json.loads(path.read_text(encoding="utf-8"))
                data[field] = value
                path.write_text(json.dumps(data), encoding="utf-8")
                with self.assertRaisesRegex(ValueError, "private_review_audit_not_qualified"):
                    inspect_private_review(folder, DAY)
        with tempfile.TemporaryDirectory() as temp:
            folder, _, _ = fixture(temp)
            path = folder / "rasd-owner-review-audit.json"
            data = json.loads(path.read_text(encoding="utf-8"))
            data["source_health"] = data["source_health"][:2]
            path.write_text(json.dumps(data), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "private_review_original_sources_not_healthy"):
                inspect_private_review(folder, DAY)

    def test_tampered_page_sha_and_missing_page_detected(self):
        with tempfile.TemporaryDirectory() as temp:
            folder, _, _ = fixture(temp)
            png = folder / "review-pages/page-001.png"
            original = png.read_bytes()
            png.write_bytes(original + b"tampered")
            with self.assertRaisesRegex(ValueError, "private_review_page_integrity_mismatch"):
                inspect_private_review(folder, DAY)
            png.unlink()
            with self.assertRaisesRegex(ValueError, "private_review_page_list_mismatch"):
                inspect_private_review(folder, DAY)

    def test_tampered_manifest_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            folder, _, _ = fixture(temp)
            path = folder / "rasd-owner-review-pages-manifest.json"
            data = json.loads(path.read_text(encoding="utf-8"))
            data["pages"][0]["sha256"] = "0" * 64
            path.write_text(json.dumps(data), encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "private_review_page_integrity_mismatch"):
                inspect_private_review(folder, DAY)

    def test_non_tty_or_wrong_confirmation_never_attests(self):
        with tempfile.TemporaryDirectory() as temp:
            folder, _, _ = fixture(temp)
            details = inspect_private_review(folder, DAY)
            target = folder / "rasd-owner-human-review-attestation.json"
            with patch("sys.stdin") as stdin, patch("sys.stdout") as stdout:
                stdin.isatty.return_value = False
                stdout.isatty.return_value = True
                with self.assertRaisesRegex(ValueError, "requires_local_tty"):
                    owner_local_attestation(folder, details)
                self.assertFalse(target.exists())
                stdin.isatty.return_value = True
                with patch("builtins.input", return_value="REVIEWED 2026-10-17 99"):
                    with self.assertRaisesRegex(ValueError, "not_confirmed"):
                        owner_local_attestation(folder, details)
            self.assertFalse(target.exists())

    def test_explicit_local_interactive_confirmation_records_self_attestation_only(self):
        with tempfile.TemporaryDirectory() as temp:
            folder, _, payload = fixture(temp)
            details = inspect_private_review(folder, DAY)
            with (patch("sys.stdin") as stdin, patch("sys.stdout") as stdout,
                  patch("builtins.input", return_value="REVIEWED 2026-10-17 1")):
                stdin.isatty.return_value = True
                stdout.isatty.return_value = True
                receipt = owner_local_attestation(folder, details)
            signed = json.loads(receipt.read_text(encoding="utf-8"))
            self.assertTrue(signed["human_visual_review"])
            self.assertFalse(signed["production_delivery_approved"])
            self.assertEqual(signed["type"], "RASD_LOCAL_OWNER_SELF_ATTESTED_REVIEW_V1")
            self.assertEqual(signed["word_sha256"], details["word_sha256"])
            self.assertEqual(signed["page_sha256"], details["page_sha256"])
            self.assertNotIn("private_key", receipt.read_text())
            self.assertNotIn("ciphertext", receipt.read_text())
            with (patch("sys.stdin") as stdin, patch("sys.stdout") as stdout,
                  patch("builtins.input", return_value="REVIEWED 2026-10-17 1")):
                stdin.isatty.return_value = True
                stdout.isatty.return_value = True
                with self.assertRaisesRegex(ValueError, "existing_no_overwrite"):
                    owner_local_attestation(folder, details)
            original_audit = json.loads((folder / "rasd-owner-review-audit.json")
                                        .read_text(encoding="utf-8"))
            self.assertIs(original_audit["human_visual_review"], False)
            self.assertTrue(payload.read_text().startswith("{"))

    def test_no_owner_files_upload_or_external_workflow_created(self):
        root = Path(__file__).resolve().parents[1]
        cli = (root / "gate61_owner_private_visual_review.py").read_text()
        self.assertNotIn("requests.", cli)
        self.assertNotIn("urlopen(", cli)
        self.assertNotIn("smtplib", cli)
        self.assertNotIn("git push", cli)
        self.assertNotIn("schedule:", cli)
        self.assertIn("requires_local_tty", cli)


if __name__ == "__main__":
    unittest.main()
