"""Gate49: private visual prerequisites — legible picture size and Arabic PDF text.

Synthetic tests only. Genuine publisher Word still requires separately
verified original-page JPEG and human page-by-page visual inspection.
"""
from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document

from daily_report import render_docx
from word_render_audit import (
    _review_original_picture_dimensions,
    _verify_original_headlines_in_pdf,
)
from test_gate42_word_image_layout import fixture_image, minimal_report


class WordVisualReadabilitySafetyTests(unittest.TestCase):
    def test_standard_original_viewport_width_accepted(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "publisher-viewport.jpg"
            fixture_image(image, 1200, 900)
            output = root / "document.docx"
            render_docx(minimal_report([image]), output)
            doc = Document(output)
            _review_original_picture_dimensions(doc, expected_articles=1)
            body = "\n".join(p.text for p in doc.paragraphs)
            _verify_original_headlines_in_pdf(
                doc,
                "وزارة التخطيط تعلن متابعات اقتصادية 0\n" + body,
                1,
            )

    def test_full_page_tall_screenshot_must_not_pass_as_legible(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "too-tall.jpg"
            fixture_image(image, 1200, 5100)
            output = root / "document.docx"
            render_docx(minimal_report([image]), output)
            doc = Document(output)
            with self.assertRaisesRegex(
                ValueError, "qa_publisher_picture_too_narrow_or_outside_page"
            ):
                _review_original_picture_dimensions(doc, expected_articles=1)
            # No cropping or re-encoding the approved publisher original.
            import hashlib
            from zipfile import ZipFile
            with ZipFile(output) as archive:
                stored = [archive.read(n) for n in archive.namelist()
                          if n.startswith("word/media/")]
            self.assertEqual(len(stored), 1)
            self.assertEqual(
                hashlib.sha256(stored[0]).hexdigest(),
                hashlib.sha256(image.read_bytes()).hexdigest(),
            )

    def test_missing_pdf_arabic_original_headline_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "publisher-view.jpg"
            fixture_image(image, 1200, 900)
            output = root / "document.docx"
            render_docx(minimal_report([image]), output)
            doc = Document(output)
            for source in (" ", "إعلان اقتصادي يخص شركة أخرى وليست الوزارة"):
                with self.subTest(source=source):
                    with self.assertRaises(ValueError):
                        _verify_original_headlines_in_pdf(doc, source, 1)
            with self.assertRaisesRegex(
                ValueError, "qa_original_report_headline_count_mismatch"
            ):
                _verify_original_headlines_in_pdf(
                    doc, "وزارة التخطيط تعلن متابعات اقتصادية", 2)

    def test_missing_or_extra_pictures_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / "publisher-view.jpg"
            fixture_image(image, 1200, 900)
            output = root / "document.docx"
            render_docx(minimal_report([image]), output)
            doc = Document(output)
            with self.assertRaisesRegex(
                ValueError, "qa_original_publisher_image_count_mismatch"
            ):
                _review_original_picture_dimensions(doc, expected_articles=2)


if __name__ == "__main__":
    unittest.main()
