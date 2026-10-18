"""Gate62 local two-day evidence, real synthetic RSA/AES-GCM round-trips only.

Never treats sample reports as real publisher coverage or signs for an owner.
"""
from __future__ import annotations

import json
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from unittest.mock import patch
from zoneinfo import ZoneInfo

from docx import Document
from docx.shared import Inches

from gate43_owner_encrypted_review import pack_owner_review, unpack_owner_review
from gate61_owner_private_visual_review import (
    inspect_private_review, owner_local_attestation,
)
from gate62_two_day_owner_acceptance import (
    TYPE, VERDICT, verify_two_completed_days,
)
from test_gate43_owner_encrypted_review import _keys, audit
from test_gate50_owner_encrypted_review_pages import synthetic_private_page

DAY1 = date(2026, 10, 16)
DAY2 = date(2026, 10, 17)


def generate_real_crypto_fixture(root: Path, day: date, *, attest=True):
    root.mkdir(parents=True, exist_ok=True)
    public_key, private_key = _keys(root)
    words = root / "report"
    words.mkdir()
    pages = root / "rendered"
    sample = synthetic_private_page(pages, 1)
    doc = Document()
    doc.add_paragraph("ملف اصطناعي للاختبار بتاريخ " + day.isoformat())
    doc.add_picture(str(sample), width=Inches(3.5))
    doc.save(words / f"RASD-Unified-{day.isoformat()}.docx")
    evidence = audit("READY_QA_ONLY")
    evidence.update({
        "target_date_cairo": day.isoformat(),
        "word_created": True,
        "official_pagination_complete": True,
        "source_health": [
            {"feed_id": "hapi", "status": "ok", "verified_links": 10},
            {"feed_id": "almasryalyoum", "status": "ok", "verified_links": 20},
            {"feed_id": "mped-official", "status": "ok", "verified_links": 20},
        ],
    })
    ciphertext = root / "owner-only-encrypted.json"
    pack_owner_review(
        words, evidence, ciphertext, public_key,
        render_passed=True, review_pages_dir=pages,
    )
    recovered = root / "private-decrypted"
    unpack_owner_review(ciphertext, private_key, recovered)
    if attest:
        detail = inspect_private_review(recovered, day)
        with (patch("sys.stdin") as stdin, patch("sys.stdout") as stdout,
              patch("builtins.input",
                    return_value=f"REVIEWED {day.isoformat()} 1")):
            stdin.isatty.return_value = True
            stdout.isatty.return_value = True
            owner_local_attestation(recovered, detail)
    return recovered


class Stage2TwoDayOwnerLocalAuditTests(unittest.TestCase):
    def fixture_pair(self, root):
        return (
            generate_real_crypto_fixture(root / "day1", DAY1),
            generate_real_crypto_fixture(root / "day2", DAY2),
        )

    def test_two_distinct_completed_cairo_days_and_explicit_attestations(self):
        with tempfile.TemporaryDirectory() as temp:
            left, right = self.fixture_pair(Path(temp))
            result = verify_two_completed_days(left, DAY1, right, DAY2)
            self.assertEqual(result["status"], VERDICT)
            self.assertEqual(result["days_count"], 2)
            self.assertEqual([x["day"] for x in result["days"]],
                             [DAY1.isoformat(), DAY2.isoformat()])
            self.assertNotEqual(result["days"][0]["word_sha256"],
                                result["days"][1]["word_sha256"])
            self.assertEqual([x["article_count"] for x in result["days"]], [1, 1])
            self.assertTrue(result["local_integrity_verified"])
            self.assertTrue(result["attestations_are_self_reported_not_cryptographically_signed"])
            self.assertFalse(result["production_delivery_approved"])
            self.assertFalse(result["email_sent"])
            self.assertFalse(result["scheduler_enabled"])
            self.assertFalse(result["storage_activated"])

    def test_reversed_folder_order_still_uses_actual_source_days(self):
        with tempfile.TemporaryDirectory() as temp:
            left, right = self.fixture_pair(Path(temp))
            result = verify_two_completed_days(right, DAY2, left, DAY1)
            self.assertEqual([x["day"] for x in result["days"]],
                             [DAY1.isoformat(), DAY2.isoformat()])

    def test_duplicate_day_cannot_count_two_runs_as_two_calendar_days(self):
        with tempfile.TemporaryDirectory() as temp:
            left, right = self.fixture_pair(Path(temp))
            with self.assertRaisesRegex(ValueError, "two_distinct_cairo_days_required"):
                verify_two_completed_days(left, DAY1, right, DAY1)

    def test_missing_human_local_attestation_not_inferred_from_crypto_green(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            left = generate_real_crypto_fixture(root / "day1", DAY1, attest=False)
            right = generate_real_crypto_fixture(root / "day2", DAY2)
            with self.assertRaisesRegex(ValueError, "missing_or_symlink"):
                verify_two_completed_days(left, DAY1, right, DAY2)

    def test_tampered_attestation_sha_cannot_be_accepted(self):
        with tempfile.TemporaryDirectory() as temp:
            left, right = self.fixture_pair(Path(temp))
            path = left / "rasd-owner-human-review-attestation.json"
            receipt = json.loads(path.read_text())
            receipt["word_sha256"] = "0" * 64
            path.write_text(json.dumps(receipt))
            with self.assertRaisesRegex(ValueError, "missing_or_mismatched"):
                verify_two_completed_days(left, DAY1, right, DAY2)

    def test_tampered_real_page_invalidates_original_and_attestation(self):
        with tempfile.TemporaryDirectory() as temp:
            left, right = self.fixture_pair(Path(temp))
            p = right / "review-pages/page-001.png"
            p.write_bytes(p.read_bytes() + b"tamper")
            with self.assertRaisesRegex(ValueError, "private_review_page_integrity_mismatch"):
                verify_two_completed_days(left, DAY1, right, DAY2)

    def test_unapproved_delivery_claim_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            left, right = self.fixture_pair(Path(temp))
            p = right / "rasd-owner-human-review-attestation.json"
            receipt = json.loads(p.read_text())
            receipt["production_delivery_approved"] = True
            p.write_text(json.dumps(receipt))
            with self.assertRaisesRegex(ValueError, "missing_or_mismatched"):
                verify_two_completed_days(left, DAY1, right, DAY2)

    def test_wrong_second_report_day_cannot_reuse_first_day_artifacts(self):
        with tempfile.TemporaryDirectory() as temp:
            left, right = self.fixture_pair(Path(temp))
            with self.assertRaisesRegex(ValueError, "private_review_audit_not_qualified"):
                verify_two_completed_days(left, DAY1, right, DAY1 - timedelta(days=1))

    def test_private_attestation_future_or_pre_complete_day_timestamp_rejected(self):
        for bad in [
            "2026-10-16T00:00:01+03:00",
            "not-a-date",
            "2026-10-18T19:00:00Z",
            "2030-01-01T12:00:00+03:00",
        ]:
            with self.subTest(bad=bad), tempfile.TemporaryDirectory() as temp:
                left, right = self.fixture_pair(Path(temp))
                path = left / "rasd-owner-human-review-attestation.json"
                receipt = json.loads(path.read_text())
                receipt["reviewed_at_cairo"] = bad
                path.write_text(json.dumps(receipt))
                with self.assertRaisesRegex(ValueError, "attestation_timestamp_invalid"):
                    verify_two_completed_days(left, DAY1, right, DAY2)

    def test_existing_cloud_human_flag_false_cannot_be_overwritten(self):
        with tempfile.TemporaryDirectory() as temp:
            left, right = self.fixture_pair(Path(temp))
            for folder in (left, right):
                audit_cloud = json.loads(
                    (folder / "rasd-owner-review-audit.json").read_text())
                self.assertFalse(audit_cloud["human_visual_review"])
            verify_two_completed_days(left, DAY1, right, DAY2)

    def test_no_remote_side_effects_or_private_key_material(self):
        source = (Path(__file__).resolve().parents[1] /
                  "gate62_two_day_owner_acceptance.py").read_text()
        for forbidden in (
            "smtplib", "requests.", "urlopen(", "boto3", "schedule:",
            "secrets.", "rsa.", "serialization.load_pem_private_key",
        ):
            self.assertNotIn(forbidden, source)
        self.assertIn("production_delivery_approved", source)


if __name__ == "__main__":
    unittest.main()
