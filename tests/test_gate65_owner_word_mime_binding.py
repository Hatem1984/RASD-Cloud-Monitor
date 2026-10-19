"""Gate65: original owner-encrypted ZIP and later MIME must be same DOCX.

Only synthetic article/Word/page fixtures with ephemeral RSA keys.
No public publisher content, Gmail auth, SMTP, R2 or real owner review.
"""
from __future__ import annotations

import hashlib
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from gate43_owner_encrypted_review import pack_owner_review, unpack_owner_review
from gate61_owner_private_visual_review import inspect_private_review
from gate65_owner_bound_mime_no_send import owner_bound_stage2_mime_check
from test_gate43_owner_encrypted_review import _keys
from test_gate50_owner_encrypted_review_pages import synthetic_private_page
from test_gate63_unified_word_mime_no_send import DAY, make_report
from test_gate64_stage2_live_mime_no_send import TODAY_AFTER_FIXTURE

ROOT = Path(__file__).resolve().parents[1]
WORKFLOW = ROOT / ".github/workflows/report-unified-ministry-manual.yml"


def test_package(directory):
    root = Path(directory)
    audit, word = make_report(directory)
    pub, private = _keys(root)
    pages = root / "private-render-pages"
    synthetic_private_page(pages, 1)
    ciphertext = root / "owner-ciphertext.json"
    receipt = pack_owner_review(
        word.parent, audit, ciphertext, pub,
        render_passed=True, review_pages_dir=pages,
    )
    return audit, word, ciphertext, private, receipt


class OwnerCiphertextToMimeSameWordTests(unittest.TestCase):
    def test_cryptographically_packed_word_matches_mime_and_decrypted_review_sha(self):
        with tempfile.TemporaryDirectory() as temp:
            audit, word, ciphertext, private, receipt = test_package(temp)
            self.assertEqual(receipt["owner_packed_word_sha256"],
                             hashlib.sha256(word.read_bytes()).hexdigest())
            self.assertTrue(receipt["owner_review_pages_encrypted"])
            with patch("smtplib.SMTP_SSL") as smtp:
                result = owner_bound_stage2_mime_check(
                    audit, word.parent, owner_receipt=receipt,
                    requested_window="yesterday",
                    today_cairo=TODAY_AFTER_FIXTURE,
                )
            smtp.assert_not_called()
            self.assertEqual(result["status"], "UNIFIED_WORD_MIME_NO_SEND_QA_ONLY")
            self.assertTrue(result["mime_rehearsed"])
            self.assertFalse(result["email_sent"])
            self.assertNotIn("owner_packed_word_sha256", result)
            local = Path(temp) / "decrypted"
            unpack_owner_review(ciphertext, private, local)
            local_audit = json.loads((local / "rasd-owner-review-audit.json").read_text())
            self.assertEqual(local_audit["owner_packed_word_sha256"],
                             receipt["owner_packed_word_sha256"])
            checked = inspect_private_review(local, DAY)
            self.assertTrue(checked["owner_ciphertext_word_identity_verified"])
            self.assertEqual(checked["word_sha256"], receipt["owner_packed_word_sha256"])

    def test_mutated_docx_after_encryption_cannot_be_mime_accepted(self):
        with tempfile.TemporaryDirectory() as temp:
            audit, word, ciphertext, private, receipt = test_package(temp)
            doc = Document(word)
            doc.add_paragraph("بيانات اصطناعية إضافية بعد التغليف")
            doc.save(word)
            self.assertNotEqual(hashlib.sha256(word.read_bytes()).hexdigest(),
                                receipt["owner_packed_word_sha256"])
            with patch("gate65_owner_bound_mime_no_send.stage2_live_no_send_mime_check") as mime:
                with self.assertRaisesRegex(ValueError, "packed_word_sha256_mismatch"):
                    owner_bound_stage2_mime_check(
                        audit, word.parent, owner_receipt=receipt,
                        requested_window="yesterday", today_cairo=TODAY_AFTER_FIXTURE,
                    )
                mime.assert_not_called()
            local = Path(temp) / "decrypted"
            unpack_owner_review(ciphertext, private, local)
            self.assertEqual(hashlib.sha256(
                (local / word.name).read_bytes()).hexdigest(),
                receipt["owner_packed_word_sha256"])

    def test_rewrite_during_mime_validation_is_detected_before_return(self):
        with tempfile.TemporaryDirectory() as temp:
            audit, word, _, _, receipt = test_package(temp)

            def mutate(*a, **kw):
                doc = Document(word)
                doc.add_paragraph("تغيير اصطناعي أثناء اختبار MIME")
                doc.save(word)
                return {"status": "UNIFIED_WORD_MIME_NO_SEND_QA_ONLY",
                        "mime_rehearsed": True, "smtp_used": False,
                        "email_sent": False, "production_delivery_approved": False}

            with patch("gate65_owner_bound_mime_no_send.stage2_live_no_send_mime_check",
                       side_effect=mutate):
                with self.assertRaisesRegex(ValueError, "packed_word_sha256_mismatch"):
                    owner_bound_stage2_mime_check(
                        audit, word.parent, owner_receipt=receipt,
                        requested_window="yesterday", today_cairo=TODAY_AFTER_FIXTURE,
                    )

    def test_missing_or_wrong_packed_sha_is_not_accepted(self):
        for fake in (None, "0" * 64, "zz" * 64, True):
            with self.subTest(fake=fake), tempfile.TemporaryDirectory() as temp:
                audit, word, _, _, receipt = test_package(temp)
                receipt["owner_packed_word_sha256"] = fake
                with self.assertRaisesRegex(ValueError, "packed_word_sha"):
                    owner_bound_stage2_mime_check(
                        audit, word.parent, owner_receipt=receipt,
                        requested_window="yesterday", today_cairo=TODAY_AFTER_FIXTURE,
                    )

    def test_hold_archive_has_no_sha_and_mime_skips(self):
        with tempfile.TemporaryDirectory() as temp:
            audit, word, _, _, _ = test_package(temp)
            word.unlink()
            audit.update(status="HOLD", coverage="DEGRADED",
                         word_created=False, all_sources_healthy=False)
            root = Path(temp)
            pub, _ = _keys(root)
            receipt = pack_owner_review(word.parent, audit, root / "hold-only.json", pub)
            self.assertIsNone(receipt["owner_packed_word_sha256"])
            result = owner_bound_stage2_mime_check(
                audit, word.parent, owner_receipt=receipt,
                requested_window="yesterday", today_cairo=TODAY_AFTER_FIXTURE,
            )
            self.assertEqual(result["status"], "SKIPPED_SOURCE_OR_PROOF_HOLD")
            self.assertFalse(result["mime_rehearsed"])
            self.assertFalse(result["email_sent"])

    def test_decrypted_owner_audit_tampered_sha_fails_local_review(self):
        with tempfile.TemporaryDirectory() as temp:
            _, _, ciphertext, private, _ = test_package(temp)
            folder = Path(temp) / "owner"
            unpack_owner_review(ciphertext, private, folder)
            target = folder / "rasd-owner-review-audit.json"
            data = json.loads(target.read_text())
            data["owner_packed_word_sha256"] = "0" * 64
            target.write_text(json.dumps(data))
            with self.assertRaisesRegex(ValueError, "owner_packed_word_sha256_mismatch"):
                inspect_private_review(folder, DAY)

    def test_legacy_real_owner_audit_without_new_field_can_still_be_reviewed(self):
        with tempfile.TemporaryDirectory() as temp:
            _, _, ciphertext, private, _ = test_package(temp)
            folder = Path(temp) / "owner"
            unpack_owner_review(ciphertext, private, folder)
            target = folder / "rasd-owner-review-audit.json"
            data = json.loads(target.read_text())
            data.pop("owner_packed_word_sha256")
            target.write_text(json.dumps(data))
            checked = inspect_private_review(folder, DAY)
            self.assertFalse(checked["owner_ciphertext_word_identity_verified"])
            self.assertFalse(checked["human_visual_review"])

    def test_owner_ciphertext_receipt_claims_must_be_original_and_private(self):
        for bad in (
            {"encrypted": False},
            {"plaintext_artifact": True},
            {"human_visual_review": True},
            {"email_sent": True},
        ):
            with self.subTest(bad=bad), tempfile.TemporaryDirectory() as temp:
                audit, word, _, _, receipt = test_package(temp)
                receipt.update(bad)
                with self.assertRaisesRegex(ValueError, "owner_encrypted_receipt_invalid"):
                    owner_bound_stage2_mime_check(
                        audit, word.parent, owner_receipt=receipt,
                        requested_window="yesterday", today_cairo=TODAY_AFTER_FIXTURE,
                    )

    def test_live_workflow_binds_same_owner_encrypted_word_after_pack(self):
        content = WORKFLOW.read_text(encoding="utf-8")
        self.assertIn("workflow_dispatch:", content)
        self.assertNotIn("schedule:", content)
        self.assertIn("github.actor == 'Hatem1984'", content)
        self.assertIn("from gate65_owner_bound_mime_no_send import", content)
        self.assertIn("owner_receipt=receipt", content)
        self.assertLess(content.index("receipt = pack_owner_review("),
                        content.index("owner_bound_stage2_mime_check("))
        self.assertIn("if: ${{ always() }}", content)
        self.assertIn("retention-days: 1", content)
        self.assertNotIn("RASD_GMAIL_APP_PASSWORD", content)
        self.assertIn("RASD_UNIFIED_WORD_MIME_NO_SEND_STATUS", content)
        self.assertIn("OWNER_CIPHERTEXT_PRESERVED", content)

    def test_gate65_never_exposes_sha_plaintext_or_smtp(self):
        content = (ROOT / "gate65_owner_bound_mime_no_send.py").read_text()
        for forbidden in ("gmail_submit(", "smtplib", "urlopen(", "requests.",
                          "PrivateGitHubLedger", "boto3", "create_bucket",
                          "send_message(", "RASD_GMAIL_APP_PASSWORD"):
            self.assertNotIn(forbidden, content)
        self.assertNotIn('"owner_packed_word_sha256": expected', content)


if __name__ == "__main__":
    unittest.main()
