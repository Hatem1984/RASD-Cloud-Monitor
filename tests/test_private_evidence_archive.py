"""Fake-S3 offline integrity checks; never contact an actual storage service."""
from __future__ import annotations

import io
import json
import tempfile
import warnings
import zipfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document

from private_evidence_archive import PrivateEvidenceArchive
from private_evidence_restore import restore_verified_bundle
from hashlib import sha256
from unittest.mock import patch


class Precondition(Exception):
    def __init__(self):
        self.response = {"Error": {"Code": "PreconditionFailed"}}


class FakeS3:
    def __init__(self):
        self.objects = {}
        self.corrupt_reads = False

    def put_object(self, *, Bucket, Key, Body, ContentType, Metadata, IfNoneMatch):
        assert IfNoneMatch == "*"
        assert Bucket == "rasd-private-evidence"
        assert "acl" not in str(Metadata).lower()
        if Key in self.objects:
            raise Precondition()
        self.objects[Key] = bytes(Body)

    def get_object(self, *, Bucket, Key):
        raw = self.objects[Key]
        if self.corrupt_reads:
            raw = raw[:-1] + bytes([raw[-1] ^ 1])
        return {"ContentLength": len(raw), "Body": io.BytesIO(raw)}


def screenshot(path: Path):
    pic = Image.new("RGB", (1200, 900), "white")
    paint = ImageDraw.Draw(pic)
    paint.rectangle((0, 0, 1199, 220), fill=(16, 30, 56))
    for y in range(235, 850, 25):
        paint.rectangle((32, y, 1100, y+14), fill=(45, 60, 75))
    pic.save(path, quality=90)


def ready_report():
    return {"status": "READY", "coverage": "HEALTHY",
            "articles": 1, "screenshots_missing": 0,
            "word_created": True}


class PrivateArchiveTests(unittest.TestCase):
    def setUp(self):
        self.client = FakeS3()
        self.store = PrivateEvidenceArchive(self.client, bucket="rasd-private-evidence")

    def proof(self, root: Path):
        image = root / "screenshot.jpg"
        screenshot(image)
        doc = Document()
        doc.add_paragraph("Verified original publisher")
        doc.add_picture(str(image))
        word = root / "RASD-Daily-2026-10-14.docx"
        doc.save(word)
        key = "a" * 24
        digest = sha256(image.read_bytes()).hexdigest()
        return word, [(key, image, digest)]

    def test_original_jpeg_and_word_are_privately_archived_and_read_back(self):
        with tempfile.TemporaryDirectory() as work:
            word, images = self.proof(Path(work))
            result = self.store.publish_ready(report=ready_report(), word=word, evidence=images)
            self.assertEqual(result["status"], "PRIVATE_VERIFIED")
            self.assertEqual(result["screenshots_archived"], 1)
            self.assertTrue(result["manifest_key"].endswith(".json"))
            self.assertEqual(len(self.client.objects), 3)
            # A rerun must be byte-idempotent, not overwrite existing proof.
            same = self.store.publish_ready(report=ready_report(), word=word, evidence=images)
            self.assertEqual(same["manifest_key"], result["manifest_key"])
            self.assertEqual(len(self.client.objects), 3)

    def test_swapped_original_article_images_fail_before_first_remote_put(self):
        """Genuine JPEGs attached to wrong articles must never be archived."""
        with tempfile.TemporaryDirectory() as work:
            folder = Path(work)
            word, evidence = self.proof(folder)
            second = folder / "second.jpg"
            screenshot(second)
            with Image.open(second) as picture:
                paint = ImageDraw.Draw(picture)
                paint.rectangle((200, 300, 600, 650), fill=(240, 20, 30))
                picture.save(second, quality=90)
            second_proof = ("b" * 24, second, sha256(second.read_bytes()).hexdigest())
            # The Word contains exactly the correct JPEG byte set, but the
            # pictures are reversed relative to the independently captured
            # article evidence order.
            doc = Document()
            doc.add_paragraph("Article one")
            doc.add_picture(str(second))
            doc.add_paragraph("Article two")
            doc.add_picture(str(evidence[0][1]))
            doc.save(word)
            report = ready_report()
            report["articles"] = 2
            with self.assertRaisesRegex(
                    RuntimeError, "private_archive_embedded_screenshot_mismatch"):
                self.store.publish_ready(
                    report=report, word=word, evidence=evidence + [second_proof])
        self.assertFalse(self.client.objects)

    def test_duplicate_docx_media_zip_member_fails_before_first_remote_put(self):
        with tempfile.TemporaryDirectory() as work:
            word, evidence = self.proof(Path(work))
            with zipfile.ZipFile(word) as package:
                image = next(name for name in package.namelist()
                             if name.startswith("word/media/"))
                payload = package.read(image)
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)
                with zipfile.ZipFile(word, "a") as package:
                    package.writestr(image, payload)
            with self.assertRaisesRegex(
                    RuntimeError, "private_archive_embedded_screenshot_mismatch"):
                self.store.publish_ready(
                    report=ready_report(), word=word, evidence=evidence)
        self.assertFalse(self.client.objects)

    def test_offline_private_restore_rechecks_original_word_and_jpeg_bytes(self):
        with tempfile.TemporaryDirectory() as work:
            word, evidence = self.proof(Path(work))
            receipt = self.store.publish_ready(
                report=ready_report(), word=word, evidence=evidence)
            original_object_count = len(self.client.objects)
            restored = restore_verified_bundle(
                self.store, manifest_key=receipt["manifest_key"],
                expected_manifest_sha256=receipt["manifest_sha256"],
                include_bytes=True)
            self.assertEqual(restored["status"], "PRIVATE_RESTORE_VERIFIED")
            self.assertEqual(restored["articles"], 1)
            self.assertEqual(restored["screenshots_verified"], 1)
            self.assertEqual(restored["word_bytes"], word.read_bytes())
            self.assertEqual(restored["screenshots"][0]["jpeg_bytes"],
                             evidence[0][1].read_bytes())
            self.assertEqual(len(Document(io.BytesIO(restored["word_bytes"])).inline_shapes), 1)
            self.assertEqual(len(self.client.objects), original_object_count)
            summary = restore_verified_bundle(
                self.store, manifest_key=receipt["manifest_key"],
                expected_manifest_sha256=receipt["manifest_sha256"])
            self.assertNotIn("word_bytes", summary)
            self.assertNotIn("screenshots", summary)

    def test_restore_rejects_mutated_manifest_commit_marker(self):
        with tempfile.TemporaryDirectory() as work:
            word, images = self.proof(Path(work))
            receipt = self.store.publish_ready(
                report=ready_report(), word=word, evidence=images)
            key = receipt["manifest_key"]
            raw = self.client.objects[key]
            self.client.objects[key] = raw[:-1] + b" "
            with self.assertRaisesRegex(
                    RuntimeError, "private_restore_manifest_digest_mismatch"):
                restore_verified_bundle(
                    self.store, manifest_key=key,
                    expected_manifest_sha256=receipt["manifest_sha256"])

    def test_restore_rejects_mutated_word_or_missing_screenshot(self):
        with tempfile.TemporaryDirectory() as work:
            word, images = self.proof(Path(work))
            receipt = self.store.publish_ready(
                report=ready_report(), word=word, evidence=images)
            args = {"manifest_key": receipt["manifest_key"],
                    "expected_manifest_sha256": receipt["manifest_sha256"]}
            word_key = next(k for k in self.client.objects if k.endswith(".docx"))
            original = self.client.objects[word_key]
            self.client.objects[word_key] = original[:-1] + bytes([original[-1] ^ 1])
            with self.assertRaisesRegex(
                    RuntimeError, "private_restore_object_digest_mismatch"):
                restore_verified_bundle(self.store, **args)
            self.client.objects[word_key] = original
            image_key = next(k for k in self.client.objects if k.endswith(".jpg"))
            del self.client.objects[image_key]
            with self.assertRaisesRegex(
                    RuntimeError, "private_restore_object_unavailable"):
                restore_verified_bundle(self.store, **args)

    def test_restore_rejects_unapproved_manifest_path_before_network(self):
        with self.assertRaisesRegex(
                RuntimeError, "private_restore_manifest_identity_mismatch"):
            restore_verified_bundle(
                self.store,
                manifest_key="public/2026-10-14/manifests/" + "a" * 64 + ".json",
                expected_manifest_sha256="a" * 64)
        self.assertFalse(self.client.objects)

    def test_restore_rejects_forged_manifest_path_even_with_matching_hash(self):
        with tempfile.TemporaryDirectory() as work:
            word, images = self.proof(Path(work))
            receipt = self.store.publish_ready(
                report=ready_report(), word=word, evidence=images)
            manifest = json.loads(self.client.objects[receipt["manifest_key"]])
            manifest["word"]["key"] = "../other-project/reports/not-rasd.docx"
            raw = (json.dumps(manifest, sort_keys=True, separators=(",", ":"))
                   + "\n").encode()
            digest = sha256(raw).hexdigest()
            key = "rasd/v1/2026-10-14/manifests/" + digest + ".json"
            self.client.objects[key] = raw
            with self.assertRaisesRegex(
                    RuntimeError, "private_restore_invalid_manifest"):
                restore_verified_bundle(
                    self.store, manifest_key=key,
                    expected_manifest_sha256=digest)

    def test_hold_cannot_upload_any_object(self):
        with tempfile.TemporaryDirectory() as work:
            word, images = self.proof(Path(work))
            report = ready_report()
            report["status"] = "HOLD"
            with self.assertRaisesRegex(RuntimeError, "private_archive_requires_complete_report"):
                self.store.publish_ready(report=report, word=word, evidence=images)
        self.assertFalse(self.client.objects)

    def test_incorrect_image_hash_cannot_upload(self):
        with tempfile.TemporaryDirectory() as work:
            word, images = self.proof(Path(work))
            images = [(images[0][0], images[0][1], "0" * 64)]
            with self.assertRaisesRegex(RuntimeError, "private_archive_screenshot_mismatch"):
                self.store.publish_ready(report=ready_report(), word=word, evidence=images)
        self.assertFalse(self.client.objects)

    def test_remote_data_corruption_fails_no_commit_manifest(self):
        with tempfile.TemporaryDirectory() as work:
            word, images = self.proof(Path(work))
            self.client.corrupt_reads = True
            with self.assertRaisesRegex(RuntimeError, "private_evidence_digest_readback_mismatch"):
                self.store.publish_ready(report=ready_report(), word=word, evidence=images)
        self.assertFalse(any("/manifests/" in key for key in self.client.objects))

    def test_archive_refuses_word_missing_the_approved_screenshot(self):
        with tempfile.TemporaryDirectory() as work:
            root = Path(work)
            word, images = self.proof(root)
            other_doc = Document()
            other_doc.add_paragraph("Missing screenshot despite a publisher proof.")
            other_doc.save(word)
            with self.assertRaisesRegex(RuntimeError, "private_archive_embedded_screenshot_mismatch"):
                self.store.publish_ready(report=ready_report(), word=word, evidence=images)
        self.assertFalse(self.client.objects)

    def test_later_invalid_screenshot_uploads_nothing(self):
        with tempfile.TemporaryDirectory() as work:
            root = Path(work)
            word, images = self.proof(root)
            bad = root / "bad.jpg"
            bad.write_bytes(b"BAD BYTES" * 1000)
            image2 = ("b" * 24, bad, "0" * 64)
            report = ready_report()
            report["articles"] = 2
            with self.assertRaisesRegex(RuntimeError, "private_archive_screenshot_mismatch"):
                self.store.publish_ready(report=report, word=word,
                                         evidence=images + [image2])
        self.assertFalse(self.client.objects)

    def test_batch_budget_refuses_before_any_upload(self):
        with tempfile.TemporaryDirectory() as work:
            word, images = self.proof(Path(work))
            with patch("private_evidence_archive.MAX_BATCH_BYTES", 9000):
                with self.assertRaisesRegex(ValueError, "private_archive_batch_budget_exceeded"):
                    self.store.publish_ready(report=ready_report(), word=word,
                                             evidence=images)
        self.assertFalse(self.client.objects)

    def test_illegal_bucket_and_prefix_blocked(self):
        with self.assertRaises(ValueError):
            PrivateEvidenceArchive(self.client, bucket="../public")
        with self.assertRaises(ValueError):
            PrivateEvidenceArchive(self.client, bucket="rasd-private-evidence", prefix="public")


if __name__ == "__main__":
    unittest.main()
