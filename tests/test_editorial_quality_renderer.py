"""QA-only Word renderer contract with no email, network or fake evidence."""
import tempfile,unittest
from pathlib import Path
from docx import Document
from daily_report import render_docx

DAY="2026-10-24"
def sample():
 return {
  "title":"مَآل | MA'AL — تجربة تحريرية",
  "display_window_start":DAY,"display_window_end":DAY,
  "coverage":"HEALTHY","verified_links":1,"sources":[],
  "preview_only":True,"window_article_count":1,"truncated":False,
  "editorial_archive_only":True,"editorial_allowed_hosts":
    ["www.mped.gov.eg"],
  "limitations":"عينة تاريخية اختبارية، لا تعني اكتمال الرصد أو إرسال البريد.",
  "articles":[{
   "headline":"وزير التخطيط: نستهدف جذب رأس مال مخاطر بـ 5 مليارات دولار للشركات الناشئة",
   "publisher_excerpt":"","source":"وزارة التخطيط — الموقع الرسمي",
   "url":"https://www.mped.gov.eg/singlenews?id=7275&lang=ar",
   "published_at_cairo":DAY+" (date only)",
   "published_display_cairo":"24 أكتوبر 2026 — يوم النشر فقط"
  }]
 }

class RenderQualityPreview(unittest.TestCase):
 def test_quality_mode_adds_evidence_sections_without_legacy_filler(self):
  with tempfile.TemporaryDirectory() as root:
   original=Path(root)/"old.docx"
   quality=Path(root)/"quality.docx"
   report=sample()
   render_docx(report,original)
   render_docx(report,quality,editorial_mode="qa_only")
   old="\n".join(p.text for p in Document(original).paragraphs)
   new="\n".join(p.text for p in Document(quality).paragraphs)
   self.assertIn("تحليل تحريري مرتبط بموضوع الخبر:",old)
   self.assertNotIn("تحليل تحريري مرتبط بموضوع الخبر:",new)
   for heading in ("الحقيقة المثبتة في العنوان:","أهم الأرقام المنشورة:",
                   "لماذا يهم الشركات؟","خطوة عملية مشروطة:","ما الذي لم يُثبت؟"):
    self.assertIn(heading,new)
   self.assertIn("5 مليارات دولار",new)
   self.assertIn("لا تمثل رصدًا حيًا",new)
   self.assertNotIn("5,000 شركة",new)
   self.assertEqual(Document(quality).core_properties.author,"MA'AL | مَآل")
 def test_qa_mode_must_be_explicit(self):
  with tempfile.TemporaryDirectory() as root:
   p=Path(root)/"x.docx"
   report=sample()
   del report["editorial_archive_only"]
   with self.assertRaisesRegex(ValueError,"requires_explicit_archive_flag"):
    render_docx(report,p,editorial_mode="qa_only")
   self.assertFalse(p.exists())
 def test_unrecognized_quality_mode_denied(self):
  with tempfile.TemporaryDirectory() as root:
   with self.assertRaisesRegex(ValueError,"editorial_mode_not_approved"):
    render_docx(sample(),Path(root)/"x.docx",editorial_mode="auto_send")
if __name__=="__main__":unittest.main()
