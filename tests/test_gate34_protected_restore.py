"""Gate 34: protected one-shot cloud restore, tested without R2 or Gmail."""
from __future__ import annotations

import io
import json
import tempfile
import unittest
from hashlib import sha256
from pathlib import Path

from docx import Document
from PIL import Image, ImageDraw

from private_evidence_archive import PrivateEvidenceArchive
from private_restore_worker import (
    EXPECTED_WORKFLOW, ONE_RUN_CONFIRMATION, OWNER_APPROVAL,
    checked_configuration, run,
)

ENV = {
    "GITHUB_EVENT_NAME": "workflow_dispatch",
    "GITHUB_REF": "refs/heads/main",
    "GITHUB_REPOSITORY": "Hatem1984/RASD-Cloud-Monitor",
    "GITHUB_WORKFLOW_REF": EXPECTED_WORKFLOW,
    "GITHUB_RUN_ATTEMPT": "1",
    "GITHUB_RUN_ID": "987654321",
    "GITHUB_ACTOR": "Hatem1984",
    "RASD_RESTORE_RUN_CONFIRMATION": ONE_RUN_CONFIRMATION,
    "RASD_R2_RESTORE_APPROVED": OWNER_APPROVAL,
    "RASD_R2_RESTORE_ACCOUNT_ID": "a" * 32,
    "RASD_R2_RESTORE_BUCKET": "rasd-private-evidence",
    "RASD_R2_RESTORE_ACCESS_KEY_ID": "test-readonly-key-not-real",
    "RASD_R2_RESTORE_SECRET_ACCESS_KEY": "test-readonly-secret-not-real",
    "RASD_R2_RESTORE_MANIFEST_DATE": "2026-10-14",
    "RASD_R2_RESTORE_MANIFEST_SHA256": "f" * 64,
}


class FakeStore:
    """Local-only mock supports initial seeding then permanently read-only."""
    def __init__(self):
        self.objects = {}
        self.read_only = False
        self.get_count = 0
        self.put_count = 0

    def put_object(self, *, Bucket, Key, Body, ContentType, Metadata, IfNoneMatch):
        if self.read_only:
            raise AssertionError("restoration MUST NEVER attempt a PUT")
        if Key in self.objects:
            raise AssertionError("unexpected idempotent duplicate in fixture")
        self.put_count += 1
        self.objects[Key] = bytes(Body)

    def get_object(self, *, Bucket, Key):
        self.get_count += 1
        data = self.objects[Key]
        return {"ContentLength": len(data), "Body": io.BytesIO(data)}


def seed_private_archive(client, directory):
    root = Path(directory)
    jpeg = root / "publisher.jpg"
    image = Image.new("RGB", (1200, 900), "white")
    paint = ImageDraw.Draw(image)
    paint.rectangle((0, 0, 1199, 225), fill=(19, 31, 77))
    for y in range(250, 890, 25):
        paint.rectangle((40, y, 1120, y + 13), fill=(43, 65, 89))
    image.save(jpeg, quality=90)
    word = root / "RASD-Daily-2026-10-14.docx"
    document = Document()
    document.add_paragraph("Original publisher evidence")
    document.add_picture(str(jpeg))
    document.save(word)
    archive = PrivateEvidenceArchive(client, bucket="rasd-private-evidence")
    return archive.publish_ready(
        report={"status": "READY", "coverage": "HEALTHY",
                "articles": 1, "screenshots_missing": 0,
                "word_created": True},
        word=word,
        evidence=[("a" * 24, jpeg, sha256(jpeg.read_bytes()).hexdigest())],
    )


class Gate34ReadOnlyRecoveryTests(unittest.TestCase):
    def test_approval_missing_or_untrusted_run_stops_before_factory(self):
        overrides = (
            {"GITHUB_EVENT_NAME": "schedule"},
            {"GITHUB_REF": "refs/heads/qa"},
            {"GITHUB_REPOSITORY": "other/repo"},
            {"GITHUB_WORKFLOW_REF": "wrong"},
            {"GITHUB_RUN_ATTEMPT": "2"},
            {"GITHUB_ACTOR": "outsider"},
            {"GITHUB_RUN_ID": "not-a-number"},
            {"RASD_RESTORE_RUN_CONFIRMATION": ""},
            {"RASD_R2_RESTORE_APPROVED": ""},
        )
        for override in overrides:
            with self.subTest(override=override), self.assertRaises(RuntimeError):
                run(environ={**ENV, **override},
                    s3_factory=lambda **kwargs: self.fail("network client created"))

    def test_credentials_and_manifest_target_must_be_prevalidated(self):
        for field, value in (
            ("RASD_R2_RESTORE_ACCOUNT_ID", "bad"),
            ("RASD_R2_RESTORE_BUCKET", "../wrong"),
            ("RASD_R2_RESTORE_ACCESS_KEY_ID", ""),
            ("RASD_R2_RESTORE_SECRET_ACCESS_KEY", ""),
            ("RASD_R2_RESTORE_MANIFEST_DATE", "2026-02-30"),
            ("RASD_R2_RESTORE_MANIFEST_DATE", "../2026-10-14"),
            ("RASD_R2_RESTORE_MANIFEST_SHA256", "bad"),
        ):
            with self.subTest(field=field), self.assertRaises(RuntimeError):
                run(environ={**ENV, field: value},
                    s3_factory=lambda **kwargs: self.fail("network client created"))

    def test_full_word_and_jpeg_restore_is_get_only_and_redacted(self):
        with tempfile.TemporaryDirectory() as work:
            client = FakeStore()
            receipt = seed_private_archive(client, work)
            client.read_only = True
            previous_puts = client.put_count
            env = {**ENV, "RASD_R2_RESTORE_MANIFEST_SHA256":
                   receipt["manifest_sha256"]}
            settings = checked_configuration(env)
            self.assertEqual(settings["manifest_key"], receipt["manifest_key"])
            result = run(environ=env, s3_factory=lambda **kwargs: client)
            self.assertEqual(result["status"], "PRIVATE_RESTORE_VERIFIED")
            self.assertEqual(result["articles"], 1)
            self.assertEqual(result["screenshots_verified"], 1)
            self.assertTrue(result["read_only"])
            self.assertEqual(client.put_count, previous_puts)
            self.assertGreaterEqual(client.get_count, 6)
            output = json.dumps(result)
            for sensitive in (
                receipt["manifest_sha256"], receipt["manifest_key"],
                "rasd-private-evidence", env["RASD_R2_RESTORE_SECRET_ACCESS_KEY"],
            ):
                self.assertNotIn(sensitive, output)
            self.assertFalse(result["email_sent"])
            self.assertFalse(result["scheduler_enabled"])

    def test_corrupt_or_missing_remote_proof_refuses_success(self):
        with tempfile.TemporaryDirectory() as work:
            client = FakeStore()
            receipt = seed_private_archive(client, work)
            env = {**ENV, "RASD_R2_RESTORE_MANIFEST_SHA256":
                   receipt["manifest_sha256"]}
            key = next(k for k in client.objects if k.endswith(".jpg"))
            del client.objects[key]
            client.read_only = True
            with self.assertRaises(RuntimeError):
                run(environ=env, s3_factory=lambda **kwargs: client)

    def test_workflow_never_schedules_sends_or_uploads_public_evidence(self):
        path = Path(__file__).resolve().parent.parent / ".github/workflows/private-restore-manual.yml"
        source = path.read_text(encoding="utf-8")
        self.assertIn("workflow_dispatch:", source)
        self.assertIn("environment: rasd-evidence-private-restore-approval", source)
        self.assertIn("RASD_R2_RESTORE_SECRET_ACCESS_KEY", source)
        self.assertNotIn("schedule:", source)
        self.assertNotIn("upload-artifact@", source)
        self.assertNotIn("digest_mailer.py", source)
        self.assertNotIn("private_proof_worker.py", source)


if __name__ == "__main__":
    unittest.main()
