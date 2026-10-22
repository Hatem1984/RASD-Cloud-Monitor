"""Official-bulletin pagination regression: all original evidence survives.

Offline synthetic JPEG fixture, not an actual publisher screenshot and never
an authorization to send anything to a real recipient.
"""
import hashlib
import tempfile
import unittest
from pathlib import Path

from PIL import Image, ImageDraw
from docx import Document

from cloud_proof_report import word_images_match_capture_digests
from daily_report import render_docx
from rasd_official_daily import compact_official_word


class OfficialWordPageLayoutTests(unittest.TestCase):
    def test_compact_spacing_keeps_original_embedded_jpeg_and_arabic_copy(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            photo = root / "genuine_only_as_offline_fixture.jpg"
            bitmap = Image.new("RGB", (1200, 1100), "#f7f7f7")
            draw = ImageDraw.Draw(bitmap)
            for y in range(25, 1000, 28):
                draw.rectangle((30, y, 1130, y + 12), fill="#324c60")
            bitmap.save(photo, "JPEG", quality=87)
            proof = [{"status": "verified",
                      "sha256": hashlib.sha256(photo.read_bytes()).hexdigest()}]
            headline = "وزير التخطيط يستعرض توجهات الدولة لضبط مسار الدين الخارجي"
            source_url = "https://mped.gov.eg/singlenews?id=7273&lang=ar"
            report = {
                "title": "رصد | البيان الصحفي الرسمي للوزارة — 2026-10-22",
                "display_window_start": "2026-10-22 00:00 القاهرة",
                "display_window_end": "2026-10-22 23:59 القاهرة",
                "coverage": "OFFICIAL_SOURCE_ONLY",
                "verified_links": 1,
                "articles": [{
                    "headline": headline,
                    "source": "وزارة التخطيط والتنمية الاقتصادية — الموقع الرسمي",
                    "url": source_url,
                    "published_at_cairo": "2026-10-22",
                    "published_display_cairo": "22 أكتوبر 2026",
                    "evidence_image_path": str(photo),
                }],
                "sources": [{"feed_id": "mped-official-original-NewsList-page-1",
                             "status": "ok", "verified_links": 1,
                             "error_type": None}],
                "limitations": "هذا تقرير رسمي محدود النطاق وليس تقرير رصد صحفي شامل.",
            }
            word = root / "official.docx"
            render_docx(report, word)
            self.assertTrue(word_images_match_capture_digests(word, proof))
            compact_official_word(word)
            self.assertTrue(word_images_match_capture_digests(word, proof))
            out = Document(word)
            self.assertEqual(len(out.inline_shapes), 1)
            self.assertIn(headline, "\n".join(p.text for p in out.paragraphs))
            self.assertIn("محدود النطاق",
                          "\n".join(p.text for p in out.paragraphs))
            self.assertLessEqual(out.styles["Normal"].font.size.pt, 10)
            self.assertLessEqual(out.styles["Title"].font.size.pt, 14)
            self.assertEqual(out.sections[0].top_margin,
                             __import__("docx").shared.Inches(0.5))
            self.assertGreater(word.stat().st_size, 3_000)


if __name__ == "__main__":
    unittest.main()
