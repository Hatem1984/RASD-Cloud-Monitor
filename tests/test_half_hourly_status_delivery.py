"""Half-hour zero-news Word/SMTP policy: offline only, no real sender or ledger."""
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document
import scheduled_rasd as rasd


class HalfHourlyWordStatus(unittest.TestCase):
    def setUp(self):
        self.ledger = MagicMock()
        self.ledger.reserve.return_value = "reserved"
        self.start = datetime(2026, 10, 25, 11, 30, tzinfo=timezone.utc)
        self.now = datetime(2026, 10, 25, 11, 37, tzinfo=timezone.utc)
        self.window = "2026-10-25T11:30Z"
        self.config = {"feeds": [
            {"id": "hapi"}, {"id": "almasryalyoum"}]}
        self.health = [
            {"feed_id": "hapi", "status": "ok", "verified_links": 10},
            {"feed_id": "almasryalyoum", "status": "ok", "verified_links": 20}]
        self.dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.dir.cleanup)

    def call(self, state=None, health=None):
        return rasd.send_empty_interval_status(
            ledger=self.ledger, sender="source@gmail.com",
            recipient="review@example.com", password="offline-not-real",
            state=state if state is not None else {"entries": {}},
            config=self.config,
            health=health if health is not None else self.health,
            window=self.window, start=self.start, now=self.now,
            output=Path(self.dir.name))

    def test_true_empty_slot_sends_one_status_word_with_no_images(self):
        with patch.object(rasd, "send_one_word",
                          return_value="<rasd-test@gmail.com>") as send:
            r = self.call()
        self.assertEqual(r["status"],
                         "interval_status_smtp_accepted_unconfirmed_inbox")
        self.assertEqual(r["items"], 0)
        self.assertTrue(r["smtp_accepted"])
        args = send.call_args
        self.assertTrue(args.kwargs["subject"].startswith(
            "رصد | تقرير حالة المتابعة النصف ساعية"))
        expected = rasd.digest_key("interval-status:" + self.window)
        self.assertEqual(args.kwargs["key"], expected)
        doc = Document(args.args[0])
        self.assertEqual(len(doc.inline_shapes), 0)
        txt = " ".join(p.text for p in doc.paragraphs)
        self.assertIn("حالة المصادر", txt)
        self.assertIn("حابي", txt)
        self.assertIn("المصري اليوم", txt)
        self.assertIn("لا يثبت ذلك عدم نشر أخبار", txt)
        self.ledger.reserve.assert_called_once_with(expected, "review@example.com")
        self.ledger.mark_smtp_accepted.assert_called_once_with(
            expected, "<rasd-test@gmail.com>")

    def test_two_backup_schedulers_cannot_repeat_reserved_slot(self):
        key = rasd.digest_key("interval-status:" + self.window)
        with patch.object(rasd, "send_one_word") as send:
            r = self.call(state={"entries": {key: {"status": "reserved"}}})
        self.assertEqual(r["status"], "interval_status_already_reserved_or_sent")
        self.ledger.reserve.assert_not_called()
        send.assert_not_called()

    def test_article_accepted_same_slot_suppresses_false_empty_notice(self):
        with patch.object(rasd, "send_one_word") as send:
            r = self.call(state={"entries": {"a"*24: {
                "kind": "interval_item", "status": "smtp_accepted",
                "accepted_at": "2026-10-25T11:36:00+00:00"}}})
        self.assertEqual(r["status"], "interval_news_already_recorded_no_empty_status")
        send.assert_not_called()
        self.ledger.reserve.assert_not_called()

    def test_degraded_source_explicitly_labelled_as_partial(self):
        health = [
            {"feed_id": "hapi", "status": "ok", "verified_links": 10},
            {"feed_id": "almasryalyoum", "status": "failed", "verified_links": 0}]
        with patch.object(rasd, "send_one_word",
                          return_value="<rasd-test@gmail.com>"):
            out = self.call(health=health)
        self.assertTrue(out["source_limited"])
        doc = Document(Path(self.dir.name) /
                       "RASD-Half-Hour-No-New-Verified-News-Status.docx")
        text = " ".join(p.text for p in doc.paragraphs)
        self.assertIn("التغطية جزئية", text)
        self.assertIn("تعذّر الوصول", text)

    def test_malformed_source_health_aborts_before_mail_reservation(self):
        with patch.object(rasd, "send_one_word") as send:
            with self.assertRaisesRegex(
                RuntimeError, "source_health_metadata_missing_or_untrusted"):
                self.call(health=[{"feed_id": "hapi", "status": "ok"}])
        send.assert_not_called()
        self.ledger.reserve.assert_not_called()

    def test_publisher_photo_hold_sends_status_not_unauthenticated_news(self):
        with patch.object(rasd, "send_one_word",
                          return_value="<rasd-test@gmail.com>") as send:
            result = rasd.send_empty_interval_status(
                ledger=self.ledger, sender="source@gmail.com",
                recipient="review@example.com", password="offline-not-real",
                state={"entries": {}}, config=self.config, health=self.health,
                window=self.window, start=self.start, now=self.now,
                output=Path(self.dir.name),
                held_articles=1, waiting_articles=0)
        self.assertEqual(
            result["status"],
            "interval_evidence_hold_status_smtp_accepted_unconfirmed_inbox")
        self.assertEqual(result["held_articles"], 1)
        self.assertEqual(result["items"], 0)
        self.assertTrue(result["smtp_accepted"])
        text = " ".join(p.text for p in Document(send.call_args.args[0]).paragraphs)
        self.assertIn("تعذّر", text)
        self.assertIn("ستظل قابلة لإعادة الفحص", text)
        self.assertEqual(len(Document(send.call_args.args[0]).inline_shapes), 0)

    def test_invalid_publisher_hold_count_fails_closed(self):
        with patch.object(rasd, "send_one_word") as send:
            with self.assertRaisesRegex(RuntimeError, "interval_status_invalid_hold_count"):
                rasd.send_empty_interval_status(
                    ledger=self.ledger, sender="source@gmail.com",
                    recipient="review@example.com", password="offline-not-real",
                    state={"entries": {}}, config=self.config, health=self.health,
                    window=self.window, start=self.start, now=self.now,
                    output=Path(self.dir.name), held_articles=-1)
        send.assert_not_called()
        self.ledger.reserve.assert_not_called()

    def test_live_interval_image_hold_routes_to_status_without_article_reservation(self):
        story = {"id": "a"*24}
        with (patch.object(rasd, "require_github_scheduler"),
              patch.object(rasd, "protected_credentials",
                           return_value=("source@gmail.com", "review@example.com",
                                         "offline", self.ledger)),
              patch.object(rasd, "discover_interval_news", return_value=[story]),
              patch.object(rasd, "prepare_source_limited_verified_word",
                           side_effect=RuntimeError("no_original_publisher_page_verified")),
              patch.object(rasd, "send_empty_interval_status",
                           return_value={"status": "evidence_hold_status_test"}) as notify,
              patch.object(rasd, "reserve_stories") as reserve,
              patch.object(rasd, "scheduled_phase", return_value="interval"),
              patch.dict(rasd.os.environ, {
                  "GITHUB_EVENT_NAME": "schedule",
                  "RASD_JOB_CRON": rasd.HALF_HOUR_CRON,
              })):
            self.ledger.read.return_value = ({"entries": {}}, "sha")
            result = rasd.run_scheduled("interval", self.now)
        self.assertEqual(result["status"], "evidence_hold_status_test")
        self.assertEqual(notify.call_args.kwargs["held_articles"], 1)
        reserve.assert_not_called()

    def test_legacy_owner_once_cannot_send_new_status(self):
        from unittest.mock import patch as mock_patch
        with mock_patch.object(rasd, "require_github_scheduler"), \
             mock_patch.object(rasd, "protected_credentials",
                               return_value=("source@gmail.com", "review@example.com",
                                             "offline", self.ledger)), \
             mock_patch.object(rasd, "discover_interval_news", return_value=[]), \
             mock_patch.object(rasd, "send_empty_interval_status") as status, \
             mock_patch.object(rasd, "scheduled_phase", return_value="interval"), \
             mock_patch.dict(rasd.os.environ, {
                "GITHUB_EVENT_NAME": "push", "RASD_JOB_CRON": "",
                "RASD_OWNER_PUSH_CONFIRM": "RASD-OWNER-NEWSROOM-2026-10-24"
             }):
            self.ledger.read.return_value = ({"entries": {}}, "sha")
            result = rasd.run_scheduled(
                "interval_owner_once",
                datetime(2026, 10, 24, 14, 30, tzinfo=timezone.utc))
            self.assertEqual(result["status"], "no_new_verified_news")
        status.assert_not_called()


if __name__ == "__main__":
    unittest.main()
