"""Offline acceptance for source-limited real-evidence reports and no-duplicate mail.

Synthetic JPEG test fixtures NEVER authorize fake publisher images in production.
"""
from __future__ import annotations

import hashlib
import tempfile
import os
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import MagicMock, patch

from PIL import Image
from docx import Document

import scheduled_rasd as worker
from monitor import article_id


def make_news(suffix: str, publisher: str = "Hapi native news RSS") -> dict:
    host = "hapijournal.com" if publisher.startswith("Hapi") else "www.almasryalyoum.com"
    url = "https://" + host + "/planning-" + suffix
    return {
        "id": article_id(url), "url": url,
        "source": publisher, "headline": "وزارة التخطيط خبر " + suffix,
        "publisher_excerpt": "", "published_at_cairo": "2026-10-22 19:00 EEST",
        "published_display_cairo": "22 أكتوبر 2026",
    }


CONFIG = {"feeds": [
    {"id": "hapi", "name": "Hapi native news RSS",
     "allowed_article_hosts": ["hapijournal.com"]},
    {"id": "almasryalyoum", "name": "Al Masry Al Youm native news RSS",
     "allowed_article_hosts": ["www.almasryalyoum.com"]},
]}
HEALTHY = [{"feed_id": "hapi", "status": "ok"},
           {"feed_id": "almasryalyoum", "status": "ok"}]
DEGRADED = [{"feed_id": "hapi", "status": "ok"},
            {"feed_id": "almasryalyoum", "status": "failed"}]


def fixture_capture(article: dict, feeds: list, target: Path) -> dict:
    if article["url"].endswith("blocked"):
        return {"status": "hold", "reason": "publisher_access_refused"}
    # Offline test content only: no connection to a real publisher.
    seed = hashlib.sha256(article["id"].encode()).digest()
    Image.new("RGB", (1100, 800), (seed[0], seed[1], seed[2])).save(target, "JPEG")
    return {"status": "verified",
            "sha256": hashlib.sha256(target.read_bytes()).hexdigest()}


class SourceLimitedOriginals(unittest.TestCase):
    def setUp(self):
        # Only this offline fixture emulates a genuinely scheduled worker.
        # Production event authentication remains strictly fail-closed.
        self._scheduled_env = patch.dict(os.environ, {"GITHUB_EVENT_NAME": "schedule"})
        self._scheduled_env.start()
        self.addCleanup(self._scheduled_env.stop)

    def build(self, stories, status=HEALTHY, *, capture=fixture_capture):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        return worker.prepare_source_limited_verified_word(
            stories, CONFIG, Path(tmp.name),
            title="Offline test only", start="2026-10-22 19:00",
            end="2026-10-22 19:30", source_records=status, capture=capture)

    def test_blocked_article_does_not_bury_another_verified_original(self):
        blocked = make_news("blocked")
        good = make_news("good")
        word, accepted, rejected, limited = self.build([blocked, good])
        self.assertTrue(limited)
        self.assertEqual(rejected, 1)
        self.assertEqual([a["id"] for a in accepted], [good["id"]])
        self.assertEqual(len(Document(word).inline_shapes), 1)
        self.assertTrue(any("نطاق جزئي" in p.text for p in Document(word).paragraphs))

    def test_unavailable_feed_is_honestly_labelled_partial(self):
        good = make_news("good")
        word, accepted, rejected, limited = self.build([good], DEGRADED)
        self.assertTrue(limited)
        self.assertEqual((len(accepted), rejected), (1, 0))
        text = "\n".join(p.text for p in Document(word).paragraphs)
        self.assertIn("تغطية جزئية", text)
        self.assertIn("مصادر RSS", text)

    def test_all_verified_retains_healthy_scope(self):
        word, accepted, rejected, limited = self.build([
            make_news("good-one"), make_news("good-two")])
        self.assertFalse(limited)
        self.assertEqual((len(accepted), rejected), (2, 0))
        self.assertEqual(len(Document(word).inline_shapes), 2)

    def test_all_blocked_fails_before_any_private_reservation(self):
        with self.assertRaisesRegex(RuntimeError, "no_original_publisher_page_verified"):
            self.build([make_news("blocked")])

    def test_digest_mismatch_fails_even_with_verified_receipt(self):
        def forged(article, feeds, output):
            Image.new("RGB", (1200, 800), "white").save(output, "JPEG")
            return {"status": "verified", "sha256": "0" * 64}
        with self.assertRaisesRegex(RuntimeError, "original_capture_digest_mismatch"):
            self.build([make_news("good")], capture=forged)

    def test_unapproved_article_is_rejected_before_any_browser_access(self):
        invalid = make_news("good")
        invalid["url"] = "https://external.invalid/forged"
        with self.assertRaises(ValueError):
            self.build([invalid], capture=lambda *_args: self.fail("browser called"))

    def test_feed_status_cannot_be_omitted_or_renamed(self):
        with self.assertRaisesRegex(ValueError, "source_health_metadata_missing"):
            self.build([make_news("good")], [{"feed_id": "hapi", "status": "ok"}],
                       capture=lambda *_args: self.fail("browser called"))

    def test_more_than_fifteen_cannot_silently_truncate(self):
        with self.assertRaisesRegex(ValueError, "invalid_source_limited_batch"):
            self.build([make_news(str(i)) for i in range(16)],
                       capture=lambda *_args: self.fail("browser called"))

    def test_partial_interval_reserves_and_accepts_only_verified_articles(self):
        blocked = make_news("blocked")
        good = make_news("good")
        ledger = MagicMock()
        ledger.read.return_value = ({"entries": {}}, "private-state-sha")
        records = []
        def discover(_config, _now, *, allow_degraded=False, source_records=None):
            self.assertTrue(allow_degraded)
            source_records.extend(HEALTHY)
            return [blocked, good]
        with tempfile.TemporaryDirectory() as tmp:
            doc = Path(tmp) / "test.docx"
            doc.write_bytes(b"offline-fixture")
            with patch.object(worker, "require_github_scheduler"), \
                 patch.object(worker, "scheduled_phase", return_value="interval"), \
                 patch.object(worker, "protected_credentials",
                              return_value=("sender@gmail.com", "recipient@example.org",
                                            "fixture-password", ledger)), \
                 patch.object(worker, "discover_interval_news", side_effect=discover), \
                 patch.object(worker, "prepare_source_limited_verified_word",
                              return_value=(doc, [good], 1, True)), \
                 patch.object(worker, "reserve_stories",
                              side_effect=lambda _ledger, rows, day:
                                  records.extend(rows)) as reservation, \
                 patch.object(worker, "send_one_word", return_value="<fixture-id>") as sent:
                outcome = worker.run_scheduled(
                    "interval", datetime(2026, 10, 22, 18, 30,
                                          tzinfo=timezone.utc))
        self.assertEqual([x["id"] for x in records], [good["id"]])
        self.assertEqual(outcome["items"], 1)
        self.assertTrue(outcome["source_limited"])
        self.assertEqual(outcome["original_screenshots_held"], 1)
        self.assertIn("تغطية جزئية", sent.call_args.kwargs["subject"])
        self.assertEqual(ledger.mark_smtp_accepted_many.call_count, 1)
        self.assertEqual(
            ledger.mark_smtp_accepted_many.call_args.args[0], [good["id"]])
        ledger.mark_smtp_accepted.assert_not_called()

    def test_discovery_default_still_holds_failed_feed(self):
        observed = {"sources": DEGRADED, "matched_public_articles": [],
                    "matched_article_pool_truncated": False}
        with self.assertRaisesRegex(RuntimeError, "source_not_healthy"):
            worker.discover_interval_news(
                CONFIG, datetime(2026, 10, 22, 18, 30, tzinfo=timezone.utc),
                inspector=lambda *_args, **_kwargs: observed)
        audit = []
        rows = worker.discover_interval_news(
            CONFIG, datetime(2026, 10, 22, 18, 30, tzinfo=timezone.utc),
            inspector=lambda *_args, **_kwargs: observed,
            allow_degraded=True, source_records=audit)
        self.assertEqual(rows, [])
        self.assertEqual(audit, [{**row, "verified_links": 0} for row in DEGRADED])


if __name__ == "__main__":
    unittest.main()
