"""MA'AL Gate E — ONE historical verified official signal -> CONDITIONAL business brief.

Manual, read-only, no scheduler/Gmail/private-ledger or external actions.
The historical sample is neither an up-to-date alert nor a real customer.
"""
from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path

from maal.decision_contracts import (BusinessContext, Evidence, ProposedAction, draft_decision_brief)
from monitor import article_id
from official_ministry_source import SAMPLE, canonical_official_article


@dataclass(frozen=True)
class BusinessDemoProfile:
    context: BusinessContext
    sectors_of_interest: tuple[str, ...]
    synthetic: bool = True

    def validate(self) -> None:
        self.context.validate()
        if not self.synthetic:
            raise ValueError("customer_baseline_required_not_demo")
        if not self.sectors_of_interest or any(
                tag not in ("infrastructure_finance", "renewable_energy")
                for tag in self.sectors_of_interest):
            raise ValueError("unreviewed_business_exposure")


def interpret_official_signal(receipt: dict, original_body: str, image: Path,
                              profile: BusinessDemoProfile,
                              *, checked_at: datetime | None = None) -> dict:
    """A bounded, traceable demonstration; not a financial forecast or LLM."""
    profile.validate()
    clock = checked_at or datetime.now(timezone.utc)
    if clock.tzinfo is None or clock.utcoffset() is None:
        raise ValueError("observation_time_must_be_aware")
    if not isinstance(receipt, dict) or receipt.get("status") != "verified":
        raise ValueError("original_source_verification_required")
    if receipt.get("url") != SAMPLE or canonical_official_article(receipt["url"]) != SAMPLE:
        raise ValueError("unapproved_official_sample_url")
    if receipt.get("id") != article_id(SAMPLE):
        raise ValueError("original_article_identity_mismatch")
    if receipt.get("published_date_cairo") != "2026-10-17":
        raise ValueError("original_date_only_not_verified")
    published = date.fromisoformat(receipt["published_date_cairo"])
    if published > clock.date():
        raise ValueError("future_publisher_date")
    headline = receipt.get("headline")
    if not isinstance(headline, str) or len(headline.strip()) < 20:
        raise ValueError("original_headline_invalid")
    # A topic assertion requires explicit wording from the ORIGINAL checked
    # publisher page, not a model hallucination or third-party news snippet.
    if not isinstance(original_body, str) or len(original_body) < 350:
        raise ValueError("original_body_incomplete")
    normalized = " ".join(original_body.split())
    if ("آلية ضمان" not in normalized or "البنية التحتية" not in normalized):
        raise ValueError("funding_signal_not_supported_by_original_body")
    if not isinstance(image, Path) or image.is_symlink() or not image.is_file():
        raise ValueError("missing_verified_original_screenshot")
    data = image.read_bytes()
    if not data.startswith(b"\xff\xd8") or not 8000 <= len(data) <= 3_000_000:
        raise ValueError("original_screenshot_integrity_invalid")
    proof_hash = hashlib.sha256(data).hexdigest()
    if proof_hash != receipt.get("sha256"):
        raise ValueError("original_screenshot_digest_mismatch")

    event_id = receipt["id"]
    proposals = (
        ProposedAction(
            "draft_task",
            "مراجعة أي إعلان رسمي لاحق يحدد إطلاق آلية الضمان وشروط أهلية الشركات؛ لا يُفترض فتح التقديم.",
            (event_id,),
        ),
        ProposedAction(
            "draft_report",
            "إعداد قائمة بيانات داخلية عن مشروعات الشركة المحتملة والتمويل الحالي قبل تقدير أي أثر مالي.",
            (event_id,),
        ),
    )
    # Bind the proposal to the established Gate D contract, while preserving
    # the official source's DATE-ONLY precision (not a fabricated timestamp).
    validated = draft_decision_brief(
        profile.context,
        (Evidence(
            evidence_id=event_id,
            publisher="وزارة التخطيط والتنمية الاقتصادية — المصدر الرسمي",
            original_url=SAMPLE,
            claim="مباحثات عن آلية ضمان تمويل مشروعات البنية التحتية",
            published_at=published,
            observed_at=clock,
            original_page_sha256=proof_hash,
            status="VERIFIED_ORIGINAL_PAGE",
        ),),
        proposals,
    )
    return {
        "product": "MA'AL",
        "tagline": "من الإشارة إلى القرار",
        "type": "HISTORICAL_OFFICIAL_SOURCE_DEMONSTRATION",
        "status": "DRAFT_REQUIRES_HUMAN_APPROVAL",
        "sample_company_is_synthetic": True,
        "decision_contract_validated": True,
        "facts_from_verified_inputs": validated["facts_from_verified_inputs"],
        "business_context": {
            "company_key": profile.context.company_key,
            "sector": profile.context.sector,
            "market": profile.context.market,
            "declared_exposures": list(profile.sectors_of_interest),
        },
        "evidence": {
            "evidence_id": event_id,
            "publisher": "وزارة التخطيط والتنمية الاقتصادية — المصدر الرسمي",
            "url": SAMPLE,
            "headline": headline,
            "published_day_cairo": published.isoformat(),
            "published_time": None,
            "publication_precision": "DATE_ONLY",
            "observed_at": clock.isoformat(),
            "original_page_jpeg_sha256": proof_hash,
            "source_receipt": "official_ministry_source._read_original",
            "finding": "ناقش البيان آلية لضمان تمويل مشروعات البنية التحتية؛ ولا يثبت بدء تقديم التمويل أو شروط الأهلية.",
        },
        "business_impact": {
            "type": "CONDITIONAL_OPPORTUNITY_NOT_REALIZED",
            "hypothesis": "قد يستحق الموضوع المتابعة لشركة افتراضية تعمل في البنية التحتية والطاقة المتجددة، إذا نُشرت شروط برنامج فعلي مناسب.",
            "not_established": [
                "فتح التمويل أو بدء قبول الطلبات",
                "أهلية الشركة أو استحقاقها",
                "حجم التمويل أو عائد مالي محدد",
                "تاريخ نافذ لتطبيق الآلية",
            ],
            "quantified_impact": None,
            "confidence_score": None,
            "priority": "MONITOR_OFFICIAL_NEXT_STEPS",
        },
        "proposals_not_executed": [
            {"action_type": p.kind, "description": p.description,
             "supporting_evidence_ids": list(p.supporting_evidence_ids),
             "cost_estimate": None}
            for p in proposals
        ],
        "human_approval_required": True,
        "external_actions_executed": False,
        "email_sent": False,
        "private_ledger_touched": False,
        "limitations": [
            "تجربة تاريخية وليست خبرًا عاجلًا.",
            "الشركة افتراضية والبيانات المالية غير متوفرة.",
            "تم إثبات تاريخ النشر فقط؛ لم يُنسب وقت غير منشور للمصدر.",
            "التحليل الحالي قائم على قواعد شفافة وليس استدلال نموذج لغوي آلي.",
        ],
    }


def write_demo_docx(brief: dict, screenshot: Path, output: Path) -> None:
    """Standalone Word proof; not sent or published; image hash is checked."""
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if brief.get("status") != "DRAFT_REQUIRES_HUMAN_APPROVAL":
        raise ValueError("only_approved_demo_shape_may_render")
    output.parent.mkdir(parents=True, exist_ok=True)
    document = Document()
    section = document.sections[0]
    section.header_distance = Inches(0.45)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    title = document.add_heading("مَآل | MA'AL — من الإشارة إلى القرار", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    document.add_paragraph("تجربة تاريخية توضيحية — شركة افتراضية — لا تمثل تقريرًا إنتاجيًا")
    ev = brief["evidence"]
    document.add_heading("1 | الإشارة الموثّقة", level=2)
    document.add_paragraph(ev["headline"])
    document.add_paragraph("المصدر: " + ev["publisher"])
    document.add_paragraph("تاريخ النشر (بدون ساعة): " + ev["published_day_cairo"])
    document.add_paragraph("الرابط: " + ev["url"])
    document.add_paragraph("الحقيقة المسندة: " + ev["finding"])
    document.add_picture(str(screenshot), width=Inches(5.7))
    document.add_paragraph("بصمة صورة الناشر SHA-256: " + ev["original_page_jpeg_sha256"])
    document.add_heading("2 | سياق النشاط", level=2)
    business = brief["business_context"]
    document.add_paragraph(business["sector"] + " — " + business["market"])
    document.add_paragraph("ملف العميل افتراضي؛ لا توجد بيانات مالية أو تشغيلية معتمدة.")
    document.add_heading("3 | الأثر المحتمل", level=2)
    impact = brief["business_impact"]
    document.add_paragraph(impact["hypothesis"])
    for item in impact["not_established"]:
        document.add_paragraph("غير مثبت: " + item, style="List Bullet")
    document.add_heading("4 | توصيات تحتاج موافقة", level=2)
    for item in brief["proposals_not_executed"]:
        document.add_paragraph(item["description"], style="List Bullet")
    document.add_heading("5 | حدود النتيجة", level=2)
    for line in brief["limitations"]:
        document.add_paragraph(line, style="List Bullet")
    document.add_paragraph("لم يُرسل بريد ولم ينفذ إجراء خارجي أو تعديل سجل تسليم.")
    document.save(output)
    # Guarantee EXACT image bytes are bound into one DOCX media object.
    from cloud_proof_report import word_images_match_capture_digests
    if not word_images_match_capture_digests(output, [
            {"status": "verified", "sha256": ev["original_page_jpeg_sha256"]}
    ]):
        output.unlink(missing_ok=True)
        raise ValueError("docx_image_does_not_match_original")
