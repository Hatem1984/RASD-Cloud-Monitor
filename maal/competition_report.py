"""Source-hash-bound Arabic briefing generated from ONE current model receipt.

No public reporter or new evidence collector. Output is explicitly a historical
AI proposal pending human editorial and owner approval.
"""
from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from maal.competition_journey import assemble_agent_journey


def _rtl(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    props=p._p.get_or_add_pPr()
    bidi=OxmlElement("w:bidi")
    bidi.set(qn("w:val"),"1")
    props.append(bidi)
    return p


def _p(doc, text, bold=False, size=11, color=None):
    paragraph=_rtl(doc.add_paragraph())
    run=paragraph.add_run(str(text))
    run.bold=bold
    run.font.name="Arial"
    run.font.size=Pt(size)
    if color: run.font.color.rgb=RGBColor(*color)
    return paragraph


def create_agent_brief_docx(case, ai, original_jpeg):
    """No fake images/facts. Verifies source binary and model binding first."""
    journey=assemble_agent_journey(case,ai)
    if (not isinstance(original_jpeg,bytes) or
            sha256(original_jpeg).hexdigest()!=case["evidence_digest"]):
        raise ValueError("publisher_original_jpeg_sha256_mismatch")
    try:
        from PIL import Image
        image=Image.open(BytesIO(original_jpeg))
        image.verify()
    except Exception as exc:
        raise ValueError("unreadable_original_image") from exc
    doc=Document()
    section=doc.sections[0]
    section.top_margin=Inches(.65)
    section.bottom_margin=Inches(.6)
    section.left_margin=Inches(.75)
    section.right_margin=Inches(.75)
    title=_p(doc,"مَآل | MA'AL — من الإشارة إلى القرار",True,21,(8,37,56))
    title.space_after=Pt(8)
    _p(doc,"ملف تحليل أعمال مولّد من استدلال Qwen3 فعلي · مسودة للمراجعة البشرية",True,11,(0,112,119))
    _p(doc,"تجربة على مصدر حكومي تاريخي وشركة افتراضية. ليس رصدًا حيًا ولا تقريرًا معتمدًا.",False,10,(130,67,30))
    _p(doc,"تاريخ نشر المصدر: "+case["published_day"]+" (اليوم فقط؛ ساعة النشر غير مثبتة)")
    _p(doc,"الناشر: "+case["publisher"])
    _p(doc,"معرف الدليل: "+case["evidence_id"])
    _p(doc,"بصمة صفحة الناشر الأصلية SHA-256: "+case["evidence_digest"],False,8)
    _p(doc,"معرف تشغيل النموذج: "+ai["run_id"],False,9)
    _p(doc,"تاريخ حفظ الاستدلال: "+ai["recorded_at"],False,9)

    def section_text(label, content):
        _p(doc,label,True,13,(8,37,56))
        if isinstance(content,(list,tuple)):
            for item in content:
                _p(doc,"• "+str(item))
        else:
            _p(doc,content)
    section_text("ما الذي حدث؟",case["headline"])
    section_text("الواقعة المقيدة بالدليل الرسمي",ai["facts"])
    section_text("لماذا قد يهم النشاط؟ — أثر مشروط",ai["inference"])
    section_text("البدائل المطروحة للمراجعة",ai["alternatives"])
    section_text("التوصية غير المعتمدة",ai["recommendation"])
    section_text("ما الذي لم يثبت؟",ai["unknowns"])
    section_text("حدود السلطة والتنفيذ",
        "لا يملك النموذج صلاحية اعتماد قرار أو إنشاء مهمة. التنفيذ الداخلي يتطلب تحققًا مستقلًا من حساب المالك، ولا توجد إجراءات خارجية.")
    section_text("المصدر الأصلي",case["original_url"])
    _p(doc,"صورة الصفحة الأصلية — تم التحقق من مطابقة بايتاتها للبصمة المذكورة",True,11,(0,112,119))
    photo=_rtl(doc.add_paragraph())
    photo.add_run().add_picture(BytesIO(original_jpeg),width=Inches(5.9))
    _p(doc,"ملاحظة تحريرية: توصيات النموذج تخضع لمراجعة بشرية. لم تثبت أهلية تمويل أو آثار مالية أو إجراءات منفذة.",False,9,(130,67,30))
    doc.core_properties.title="MAAL Source-Linked Qwen3 Business Review (Historical QA)"
    doc.core_properties.subject="Unapproved AI business decision proposal linked to verified source image"
    output=BytesIO()
    doc.save(output)
    binary=output.getvalue()
    with ZipFile(BytesIO(binary)) as archive:
        media=[x for x in archive.namelist() if x.startswith("word/media/") and not x.endswith("/")]
        if len(media)!=1 or sha256(archive.read(media[0])).hexdigest()!=case["evidence_digest"]:
            raise ValueError("docx_original_image_digest_mismatch")
    if (journey["ai"]["run_id"] not in
            " ".join(p.text for p in Document(BytesIO(binary)).paragraphs)):
        raise ValueError("docx_model_run_identity_missing")
    return binary
